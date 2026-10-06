# -*- coding: utf-8 -*-
"""
METALIC ERP COSTING & MRP AUTOMATION PORTAL
===========================================
Ứng dụng Streamlit tự động hóa:
  • Rã BOM Kỹ thuật đa cấp (Technical BOM) -> BOM ERP (FG / SG / BTP / RM)
  • Hoạch định nhu cầu vật tư MRP đệ quy (R_cum nhân dồn theo mọi đường đi)
  • Hạch toán giá thành TK 621 / 622 / 627 và Báo cáo P&L dự án
  • Kết xuất Excel Master 25 sheet với 100% công thức liên kết sống

Chạy:  streamlit run app.py
"""
import csv
import datetime as dt
import hashlib
import io
import math
import re
import sys
import unicodedata
from collections import defaultdict

import openpyxl
import pandas as pd
import streamlit as st
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.setrecursionlimit(10000)

APP_TITLE = "METALIC ERP COSTING & MRP AUTOMATION PORTAL"

# =============================================================================
# 1. TIỆN ÍCH CHUẨN HÓA DỮ LIỆU
# =============================================================================


def is_blank(v):
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    if isinstance(v, str):
        return v.strip() == ""
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def norm(s):
    """Chuẩn hóa chuỗi: bỏ dấu tiếng Việt, chữ thường, gộp khoảng trắng."""
    if is_blank(s):
        return ""
    s = str(s).replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = s.lower().replace("\n", " ").replace("\r", " ")
    return re.sub(r"\s+", " ", s).strip()


def to_str(v):
    if is_blank(v):
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def to_float(v, default=0.0):
    if is_blank(v):
        return default
    if isinstance(v, bool):
        return float(v)
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "")
    s = re.sub(r"[^\d,.\-eE]", "", s)
    s = re.sub(r"[eE](?!\d|-\d)", "", s)
    if not s or s in ("-", ".", ","):
        return default
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        if re.fullmatch(r"-?\d{1,3}(,\d{3})+", s):
            s = s.replace(",", "")
        else:
            s = s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return default


def ncode(v):
    """Chuẩn hóa mã để so khớp: bỏ khoảng trắng, viết hoa."""
    return re.sub(r"\s+", "", to_str(v)).upper()


def slug(s, sep="_"):
    n = norm(s).upper()
    return re.sub(r"[^A-Z0-9.]+", sep, n).strip(sep)


def fmt_dim(v):
    if v is None:
        return ""
    if float(v).is_integer():
        return str(int(v))
    return f"{v:g}"


def fmt_plate(v):
    if abs(v * 10 - round(v * 10)) < 1e-9:
        return f"{v:.1f}"
    return fmt_dim(v)


def truthy_mark(v):
    """Ô công đoạn được đánh dấu? (x, X, v, ✓, 1, có, số > 0 ...)."""
    if is_blank(v) or v is False:
        return False
    if v is True:
        return True
    if isinstance(v, (int, float)):
        return v != 0
    s = norm(v)
    return s not in ("", "0", "-", "false", "no", "khong", "n", "none", "nan", "0.0")


def file_hash(data):
    return hashlib.md5(data).hexdigest()[:12]


# =============================================================================
# 2. HÀM BẮT BUỘC XỬ LÝ Ô GỘP (MergedCell) – THEO ĐẶC TẢ MỤC 6
# =============================================================================


def safe_set_cell(ws, row_or_coord, col=None, value=None):
    try:
        cell = ws[row_or_coord] if isinstance(row_or_coord, str) else ws.cell(row=row_or_coord, column=col)
        if type(cell).__name__ == "MergedCell":
            for rng in ws.merged_cells.ranges:
                if cell.coordinate in rng:
                    top_left = ws.cell(row=rng.min_row, column=rng.min_col)
                    top_left.value = value
                    return top_left
            return None
        cell.value = value
        return cell
    except Exception:
        return None


# =============================================================================
# 3. ĐỌC FILE THÔ (XLSX/CSV) -> LƯỚI GIÁ TRỊ (có xử lý ô gộp & khoảng lùi)
# =============================================================================

MAX_ROWS = 20000
MAX_COLS = 150


def _read_csv_grid(data):
    text = None
    for enc in ("utf-8-sig", "cp1258", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))[:MAX_ROWS]
    ncols = min(max((len(r) for r in rows), default=0), MAX_COLS)
    raw = [[(r[c] if c < len(r) and r[c] != "" else None) for c in range(ncols)] for r in rows]
    indent = [[(len(v) - len(v.lstrip(" ")) if isinstance(v, str) else 0) for v in r] for r in raw]
    pct = [[False] * len(r) for r in raw]
    return [{"name": "CSV", "raw": raw, "filled": [list(r) for r in raw], "indent": indent, "pct": pct}]


def read_grids(data, filename):
    """Trả về danh sách sheet: raw (chỉ ô thực), filled (lan giá trị ô gộp), indent."""
    if filename.lower().endswith(".csv"):
        return _read_csv_grid(data)
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        max_r = min(ws.max_row or 0, MAX_ROWS)
        max_c = min(ws.max_column or 0, MAX_COLS)
        if max_r == 0 or max_c == 0:
            continue
        raw = [[None] * max_c for _ in range(max_r)]
        indent = [[0] * max_c for _ in range(max_r)]
        pct = [[False] * max_c for _ in range(max_r)]
        for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c)):
            for c_idx, cell in enumerate(row):
                if type(cell).__name__ == "MergedCell":
                    continue
                v = cell.value
                raw[r_idx][c_idx] = v
                pct[r_idx][c_idx] = "%" in str(cell.number_format or "")
                ind = 0
                try:
                    ind = int((cell.alignment.indent or 0) * 4)
                except Exception:
                    ind = 0
                if isinstance(v, str):
                    ind += len(v) - len(v.lstrip(" 　"))
                indent[r_idx][c_idx] = ind
        filled = [list(r) for r in raw]
        for rng in ws.merged_cells.ranges:
            r0, c0 = rng.min_row - 1, rng.min_col - 1
            if r0 >= max_r or c0 >= max_c:
                continue
            v = raw[r0][c0]
            for rr in range(r0, min(rng.max_row, max_r)):
                for cc in range(c0, min(rng.max_col, max_c)):
                    filled[rr][cc] = v
                    indent[rr][cc] = indent[r0][c0]
        # bỏ các dòng trống cuối
        while filled and all(is_blank(v) for v in filled[-1]):
            filled.pop()
            raw.pop()
            indent.pop()
            pct.pop()
        out.append({"name": ws.title, "raw": raw, "filled": filled, "indent": indent, "pct": pct})
    return out


# =============================================================================
# 4. SMART HEADER PARSER
# =============================================================================

# (field, [regex ưu tiên], regex loại trừ) – so khớp trên tiêu đề đã chuẩn hóa
SO_SPEC = [
    ("Mã Sản Phẩm ERP", [r"\bma\b.*\berp\b", r"\berp\b.*\b(code|ma)\b", r"^erp\b"], None),
    ("Mã Sản Phẩm KH", [r"\bma\b.*\b(kh|khach hang|khach)\b", r"customer.*(code|item|part|no)", r"\b(cust|buyer) (item|code)\b"], None),
    ("Giá Bán FOB ($USD)", [r"\bfob\b", r"gia ban", r"unit price|don gia|\bprice\b|\busd\b"], r"thanh tien|amount|total|tong"),
    ("Số Lượng", [r"^(so luong|sl|qty|quantity)\b", r"so luong|\bsl\b|\bqty\b|quantity"], r"thanh tien|amount|tong|total|pallet|thung"),
    ("Mã Sản Phẩm ERP#", [r"\bma\b.*\b(sp|san pham|tp|thanh pham|hang)\b", r"\bsku\b|fg code|product code|item code", r"^(ma|code)$"], r"\b(kh|khach)\b"),
    ("Tên Thành Phẩm", [r"\bten\b|description|mo ta|dien giai|product name"], None),
    ("Quy Cách", [r"quy cach|kich thuoc|\bspec|dimension|\bsize\b"], None),
    ("ĐVT", [r"\bdvt\b|don vi tinh|^don vi\b|^unit$|\buom\b"], None),
    ("STT", [r"^(stt|tt|no\.?|#|item)$"], None),
]

PO_SPEC = [
    ("Số PO Mua", [r"\bso po\b|\bpo (no|number|so)\b|^po$|^po\b", r"so don (dat )?(mua )?hang"], None),
    ("Nhà Cung Cấp", [r"nha cung cap|\bncc\b|supplier|vendor"], None),
    ("Mã Vật Tư ERP", [r"\bma\b.*\b(vat tu|vt|nvl|erp|hang)\b", r"material code|item code|part (no|code)", r"^(ma|code)$"], r"\bpo\b|ncc|nha cung"),
    ("Tên Hàng Hóa Phụ Kiện Mua", [r"ten hang|ten vat tu|\bten\b|description|mo ta|hang hoa"], None),
    ("Đơn Giá Mua (VND)", [r"don gia|unit price|\bgia\b|\bprice\b"], r"thanh tien|amount|tong|total"),
    ("Số Lượng Mua Thực Tế", [r"so luong|\bsl\b|\bqty\b|quantity"], r"thanh tien|amount|tong"),
    ("ĐVT", [r"\bdvt\b|don vi tinh|^don vi\b|^unit$|\buom\b"], None),
    ("Ngày Hẹn Giao", [r"ngay.*giao|hen giao|\beta\b|delivery|\bngay\b|\bdate\b"], None),
]

BOM_SPEC = [
    ("muc", [r"^muc\b", r"^(stt|tt|item|no\.?|level|cap|hang muc)$", r"^stt\b"], None),
    ("rm_code", [r"\bma\b.*\b(nvl|vat tu|vt|nguyen (vat )?lieu|phoi|vl)\b", r"material code|rm code|raw material"], r"khoi luong|\bkl\b"),
    ("fg_code", [r"\bma\b.*\b(thanh pham|tp)\b", r"\bfg code\b"], None),
    ("code", [r"\bma\b.*\b(chi tiet|ct|so|ban ve|bv|hang|erp|cum|sp|san pham|bo phan)\b",
              r"part (no|number|code)|item code|drawing no", r"^(ma|code|ky hieu|ma so)\b"],
     r"\b(nvl|vat tu|khach|kh)\b"),
    ("weight", [r"(khoi luong|trong luong|\bkl\b).*(nguyen lieu|nvl|phoi|vat lieu|\bnl\b)",
                r"raw (material )?weight|blank weight", r"khoi luong|trong luong|\bweight\b|\bkl\b"],
     r"thanh pham|\btinh\b|\bnet\b|\btong\b"),
    ("desc", [r"mo ta|ten (chi tiet|goi|hang|san pham|vat tu|cum|bo phan)|description|dien giai", r"^ten\b|\bten\b"], None),
    ("material", [r"vat lieu|^material|mac thep|^mac\b|\bgrade\b|chat lieu"], r"\bma\b|code|bill of"),
    ("thick", [r"(^|\| )(day|do day|chieu day|thk|thickness|t)( ?\(mm\)| mm)?$", r"\bdo day\b|\bchieu day\b|thickness|\bday\b"], r"\bdai\b|so luong|inch"),
    ("width", [r"(^|\| )(rong|chieu rong|width|w|phi|duong kinh|d|ø|o)( ?\(mm\)| mm)?$", r"\bchieu rong\b|\brong\b|\bwidth\b|duong kinh|\bphi\b"], r"inch|\(b\)"),
    ("width2", [r"(rong|width|canh).*\(b\)", r"\bcanh b\b|chieu cao"], r"inch"),
    ("length", [r"(^|\| )(dai|chieu dai|length|l)( ?\(mm\)| mm)?$", r"\bchieu dai\b|\bdai\b|\blength\b"], r"inch"),
    ("qty", [r"(so luong|\bsl\b).*(/|cho|tren|1)\s*(sp|bo|cum|cha|tp|chiec)", r"^(so luong|sl|qty|quantity|s\.l)\b",
             r"so luong|\bsl\b|\bqty\b|quantity"], r"\btong\b|total|\bmua\b"),
    ("uom", [r"\bdvt\b|don vi tinh|^unit$|\buom\b|^don vi\b"], None),
    ("scrap", [r"hao hut|scrap|\bhh\b|ty le hao|\bloss\b"], None),
    ("price", [r"don gia|unit price|^gia\b|\bgia (nvl|vat tu|mua|vl)\b"], r"gia cong|thanh tien"),
]

STAGE_RE = (r"\b(cat|chan|dot|dap|uon|cuon|khoan|taro|tien|phay|han|mai|lam sach|phun bi|phun cat|son|"
            r"son tinh dien|ma|ma kem|xi ma|nhung kem|lap rap|lap dat|dong goi|cnc|laser|plasma|"
            r"gia cong|nhiet luyen|danh bong|ep)\b")
STAGE_EXCL = (r"^ma (chi tiet|ct|so|ban ve|bv|hang|erp|nvl|vat|vt|tp|thanh|san|sp|kh|khach|nguyen|code|cum)|"
              r"thanh tien|so tien|tien (usd|vnd)|ghi chu|\bten\b|mo ta|khoi luong|so luong|don gia|"
              r"dien tich|\bm2\b|\btong\b|hoan thien|be mat")


def map_columns(headers, spec):
    used, mapping = set(), {}
    for field, pats, excl in spec:
        key = field.rstrip("#")
        if key in mapping:
            continue
        for p in pats:
            for ci, h in enumerate(headers):
                if ci in used or not h:
                    continue
                if excl and re.search(excl, h):
                    continue
                if re.search(p, h):
                    mapping[key] = ci
                    used.add(ci)
                    break
            if key in mapping:
                break
    return mapping


def detect_stage_columns(headers_norm, headers_disp, used):
    stages, stage_text_col = [], None
    for ci, h in enumerate(headers_norm):
        if ci in used or not h:
            continue
        parts = [p.strip() for p in h.split(" | ")]
        disp_parts = [p.strip() for p in headers_disp[ci].split(" | ")]
        sub, sub_disp = parts[-1], re.sub(r"\s+", " ", disp_parts[-1])
        if re.search(r"cong doan|quy trinh|process", h):
            if len(parts) == 1 or re.search(r"cong doan|quy trinh|process", sub):
                stage_text_col = ci if stage_text_col is None else stage_text_col
                continue
            stages.append((ci, sub_disp))
        elif re.search(STAGE_RE, sub) and not re.search(STAGE_EXCL, sub):
            stages.append((ci, sub_disp))
    return stages, stage_text_col


def _row_is_index(row):
    vals = [v for v in row if not is_blank(v)]
    if len(vals) < 4:
        return False
    nums = []
    for v in vals:
        m = re.fullmatch(r"\(?(\d+)\)?", to_str(v))
        if not m:
            return False
        nums.append(int(m.group(1)))
    return nums == sorted(nums) and len(set(nums)) == len(nums)


def _header_texts(grid, r, ncols, two):
    disp = []
    for c in range(ncols):
        top = to_str(grid[r][c]) if c < len(grid[r]) else ""
        sub = ""
        if two and r + 1 < len(grid) and c < len(grid[r + 1]):
            sub = to_str(grid[r + 1][c])
        if two and sub and top and norm(sub) != norm(top):
            disp.append(f"{top} | {sub}")
        elif two and sub and not top:
            disp.append(sub)
        else:
            disp.append(top)
    return [norm(d) for d in disp], disp


def _numeric_cells(row):
    n = 0
    for v in row:
        if isinstance(v, (int, float)) and not isinstance(v, bool) and not is_blank(v):
            n += 1
        elif isinstance(v, str) and re.fullmatch(r"[\d.,\s]+", v.strip() or "x"):
            n += 1
    return n


def detect_header(grid, spec, required, with_stages=False, scan_rows=40):
    """Tìm dòng tiêu đề (1 hoặc 2 tầng) có điểm khớp cao nhất."""
    best = None
    ncols = max((len(r) for r in grid), default=0)
    for r in range(min(scan_rows, len(grid))):
        nonempty = sum(1 for v in grid[r] if not is_blank(v))
        if nonempty < 2:
            continue
        options = [False]
        if r + 1 < len(grid):
            nxt = grid[r + 1]
            # Chỉ ghép 2 tầng khi bản thân dòng r đã là tiêu đề (>= 2 cột khớp) – tránh ghép
            # khối tiêu đề gộp phía trên (VD: "Bill of Materials" gộp H1:AD4) với dòng tiêu đề thật.
            top_only = map_columns(_header_texts(grid, r, ncols, False)[0], spec)
            if len(top_only) >= 2 and sum(1 for v in nxt if not is_blank(v)) >= 2 and _numeric_cells(nxt) <= 1:
                options.append(True)
        for two in options:
            hn, hd = _header_texts(grid, r, ncols, two)
            mp = map_columns(hn, spec)
            score = len(mp)
            stages, stxt = ([], None)
            if with_stages:
                stages, stxt = detect_stage_columns(hn, hd, set(mp.values()))
                score += len(stages) + (1 if stxt is not None else 0)
            if not all(any(k in mp for k in grp) for grp in required):
                continue
            # Hòa điểm trên cùng dòng -> ưu tiên tiêu đề 2 tầng (dòng phụ mm/inch, Dày/Rộng/Dài...)
            if best is None or score > best["score"] or (score == best["score"] and two and best["row"] == r):
                best = {"row": r, "two": two, "score": score, "map": mp, "hn": hn, "hd": hd,
                        "stages": stages, "stage_text": stxt}
    if best is None:
        return None
    # MỤC đa cột: tiêu đề "MỤC" gộp ngang nhiều cột (A5:E6) – mỗi cột là một cấp
    best["muc_cols"] = []
    if "muc" in best["map"]:
        ci = best["map"]["muc"]
        cols = [ci]
        while cols[-1] + 1 < len(best["hn"]) and best["hn"][cols[-1] + 1] == best["hn"][ci] \
                and cols[-1] + 1 not in best["map"].values():
            cols.append(cols[-1] + 1)
        best["muc_cols"] = cols
    start = best["row"] + (2 if best["two"] else 1)
    while start < len(grid) and _row_is_index(grid[start]):
        start += 1
    best["data_start"] = start
    return best


FOOTER_RE = r"^(tong cong|tong|total|cong|ghi chu|nguoi lap|nguoi duyet|duyet|phe duyet|prepared|approved|checked|ky ten)\b"


def _pick_best_sheet(grids, spec, required, with_stages=False):
    best = None
    for g in grids:
        h = detect_header(g["filled"], spec, required, with_stages)
        if h and (best is None or h["score"] > best[1]["score"]):
            best = (g, h)
    return best


# =============================================================================
# 5. PARSER SO / PO / TECHNICAL BOM
# =============================================================================

SO_COLS = ["STT", "Mã Sản Phẩm KH", "Mã Sản Phẩm ERP", "Tên Thành Phẩm", "Quy Cách", "ĐVT", "Số Lượng", "Giá Bán FOB ($USD)"]
PO_COLS = ["Số PO Mua", "Nhà Cung Cấp", "Mã Vật Tư ERP", "Tên Hàng Hóa Phụ Kiện Mua", "ĐVT",
           "Số Lượng Mua Thực Tế", "Đơn Giá Mua (VND)", "Ngày Hẹn Giao"]


@st.cache_data(show_spinner=False)
def parse_so(data, filename):
    grids = read_grids(data, filename)
    pick = _pick_best_sheet(grids, SO_SPEC, [("Số Lượng",), ("Mã Sản Phẩm ERP", "Mã Sản Phẩm KH")])
    if not pick:
        raise ValueError("Không nhận diện được tiêu đề Sales Order (cần tối thiểu cột Mã sản phẩm và Số lượng).")
    g, h = pick
    mp, grid, warns = h["map"], g["filled"], []
    if "Mã Sản Phẩm ERP" not in mp:
        warns.append("SO không có cột 'Mã Sản Phẩm ERP' – dùng 'Mã Sản Phẩm KH' làm mã ERP.")
    if "Giá Bán FOB ($USD)" not in mp:
        warns.append("SO không có cột Giá bán FOB – mặc định 0 USD.")
    rows = []
    for r in range(h["data_start"], len(grid)):
        row = grid[r]

        def gv(f):
            ci = mp.get(f)
            return row[ci] if ci is not None and ci < len(row) else None

        erp = to_str(gv("Mã Sản Phẩm ERP")) or to_str(gv("Mã Sản Phẩm KH"))
        name = to_str(gv("Tên Thành Phẩm"))
        if not erp:
            continue
        if re.search(FOOTER_RE, norm(erp)) or re.search(FOOTER_RE, norm(name)):
            continue
        qty = to_float(gv("Số Lượng"))
        if qty <= 0:
            warns.append(f"Bỏ qua dòng SO '{erp}' vì Số Lượng = {qty:g} (bắt buộc > 0).")
            continue
        rows.append({
            "STT": len(rows) + 1,
            "Mã Sản Phẩm KH": to_str(gv("Mã Sản Phẩm KH")),
            "Mã Sản Phẩm ERP": erp,
            "Tên Thành Phẩm": name,
            "Quy Cách": to_str(gv("Quy Cách")),
            "ĐVT": to_str(gv("ĐVT")) or "Cái",
            "Số Lượng": int(qty) if float(qty).is_integer() else qty,
            "Giá Bán FOB ($USD)": to_float(gv("Giá Bán FOB ($USD)")),
        })
    if not rows:
        raise ValueError("File SO không có dòng đơn hàng hợp lệ (Số Lượng > 0).")
    info = {"sheet": g["name"], "header_row": h["row"] + 1,
            "mapping": {k: h["hd"][v] for k, v in mp.items()}}
    return pd.DataFrame(rows, columns=SO_COLS), warns, info


@st.cache_data(show_spinner=False)
def parse_po(data, filename):
    grids = read_grids(data, filename)
    pick = _pick_best_sheet(grids, PO_SPEC, [("Mã Vật Tư ERP", "Tên Hàng Hóa Phụ Kiện Mua"), ("Đơn Giá Mua (VND)",)])
    if not pick:
        raise ValueError("Không nhận diện được tiêu đề Purchase Order (cần cột Mã vật tư/Tên hàng và Đơn giá).")
    g, h = pick
    mp, grid = h["map"], g["filled"]
    rows = []
    for r in range(h["data_start"], len(grid)):
        row = grid[r]

        def gv(f):
            ci = mp.get(f)
            return row[ci] if ci is not None and ci < len(row) else None

        code, name = to_str(gv("Mã Vật Tư ERP")), to_str(gv("Tên Hàng Hóa Phụ Kiện Mua"))
        if not code and not name:
            continue
        if re.search(FOOTER_RE, norm(code or name)):
            continue
        rows.append({
            "Số PO Mua": to_str(gv("Số PO Mua")),
            "Nhà Cung Cấp": to_str(gv("Nhà Cung Cấp")),
            "Mã Vật Tư ERP": code,
            "Tên Hàng Hóa Phụ Kiện Mua": name,
            "ĐVT": to_str(gv("ĐVT")),
            "Số Lượng Mua Thực Tế": to_float(gv("Số Lượng Mua Thực Tế")),
            "Đơn Giá Mua (VND)": to_float(gv("Đơn Giá Mua (VND)")),
            "Ngày Hẹn Giao": to_str(gv("Ngày Hẹn Giao")),
        })
    info = {"sheet": g["name"], "header_row": h["row"] + 1,
            "mapping": {k: h["hd"][v] for k, v in mp.items()}}
    return pd.DataFrame(rows, columns=PO_COLS), [], info


BOM_BASE_COLS = ["Sheet", "Dòng", "MỤC", "Cấp", "Mã chi tiết", "Mô tả", "Vật liệu", "Dày", "Rộng", "Rộng B", "Dài",
                 "SL/Cha", "KL NVL (kg)", "Mã NVL", "ĐVT", "Hao hụt (%)", "Đơn giá", "Mã TP"]
STAGE_PREFIX = "CĐ | "
NUM_COLS = ["Cấp", "Dày", "Rộng", "Rộng B", "Dài", "SL/Cha", "KL NVL (kg)", "Hao hụt (%)", "Đơn giá"]
DOTTED_RE = r"\d+([.\-]\d+)*\.?"


def _split_stage_text(s):
    return [p.strip() for p in re.split(r"[,;/>+\n→]|\s-\s|\s–\s", to_str(s)) if p.strip()]


@st.cache_data(show_spinner=False)
def parse_bom(data, filename):
    grids = read_grids(data, filename)
    all_rows, stage_names, sheet_ctx, infos, warns = [], {}, {}, [], []
    for g in grids:
        h = detect_header(g["filled"], BOM_SPEC, [("code", "desc", "rm_code"), ("qty", "muc", "weight")], with_stages=True)
        if not h:
            infos.append({"sheet": g["name"], "status": "Bỏ qua (không thấy tiêu đề BOM)"})
            continue
        mp, filled, raw, indent = h["map"], g["filled"], g["raw"], g["indent"]
        ctx = " ".join(to_str(v) for r in filled[: h["row"]] for v in r if not is_blank(v))
        sheet_ctx[g["name"]] = f"{g['name']} {ctx}"
        stage_cols = []
        for ci, name in h["stages"]:
            key = slug(name)
            stage_names.setdefault(key, name)
            stage_cols.append((ci, key))
        sheet_rows = []
        muc_cols = h.get("muc_cols") or []
        muc_path = []
        for r in range(h["data_start"], len(filled)):
            frow, rrow, irow = filled[r], raw[r], indent[r]

            def gv(f, use_raw=False):
                ci = mp.get(f)
                if ci is None:
                    return None
                src = rrow if use_raw else frow
                return src[ci] if ci < len(src) else None

            muc = to_str(gv("muc", True))
            if len(muc_cols) > 1:
                # MỤC đa cột: cột chứa số thứ tự cho biết cấp (A=cấp 1, B=cấp 2, ...) -> dựng chuỗi 1 / 1.1 / 1.1.1
                lvl = next((k for k, ci in enumerate(muc_cols) if ci < len(rrow) and not is_blank(rrow[ci])), None)
                if lvl is not None:
                    v = to_str(rrow[muc_cols[lvl]])
                    if re.fullmatch(r"\d+", v):
                        muc_path = (muc_path + ["0"] * lvl)[:lvl] + [v]
                        muc = ".".join(muc_path)
                    else:
                        muc = v
                else:
                    muc = ""
            code = to_str(gv("code", True))
            desc = to_str(gv("desc"))
            rmc = to_str(gv("rm_code", True))
            if not (muc or code or desc or rmc):
                continue
            if re.search(FOOTER_RE, norm(desc or code or muc)) and not code:
                continue
            ind_ci = mp.get("desc", mp.get("code"))
            ind = irow[ind_ci] if ind_ci is not None and ind_ci < len(irow) else 0
            # Công đoạn: số > 0 = thứ tự thực hiện (1, 2, 3...); dấu x/✓ = theo thứ tự cột
            stages = {}
            for pos, (ci, key) in enumerate(stage_cols, 1):
                v = frow[ci] if ci < len(frow) else None
                if truthy_mark(v):
                    num = to_float(v, 0) if isinstance(v, (int, float)) or re.fullmatch(r"\s*\d+([.,]\d+)?\s*", to_str(v)) else 0
                    stages[key] = num if num > 0 else float(pos)
            if h["stage_text"] is not None:
                for pos, nm in enumerate(_split_stage_text(frow[h["stage_text"]] if h["stage_text"] < len(frow) else None), 1):
                    key = slug(nm)
                    stage_names.setdefault(key, nm)
                    stages[key] = float(pos)
            scrap_v = gv("scrap")
            scrap = to_float(scrap_v, None) if not is_blank(scrap_v) else None
            sc_ci = mp.get("scrap")
            scrap_pct_fmt = bool(sc_ci is not None and sc_ci < len(g["pct"][r]) and g["pct"][r][sc_ci])
            if scrap is not None and scrap_pct_fmt:
                scrap *= 100  # ô định dạng % trong Excel (0.05 -> 5%)
            sheet_rows.append({
                "Sheet": g["name"], "Dòng": r + 1, "MỤC": muc, "_indent": ind,
                "Mã chi tiết": code, "Mô tả": desc, "Vật liệu": to_str(gv("material")),
                "Dày": to_float(gv("thick"), None), "Rộng": to_float(gv("width"), None),
                "Rộng B": to_float(gv("width2"), None), "Dài": to_float(gv("length"), None),
                "SL/Cha": to_float(gv("qty", True), None), "KL NVL (kg)": to_float(gv("weight", True), None),
                "Mã NVL": rmc, "ĐVT": to_str(gv("uom")), "Hao hụt (%)": scrap,
                "Đơn giá": to_float(gv("price"), None), "Mã TP": to_str(gv("fg_code")),
                "_stages": stages, "_scrap_pct_fmt": scrap_pct_fmt,
            })
        # ---- "TỈ LỆ HAO HỤT" dạng hệ số (0 < f <= 1, VD = 1) thay vì phần trăm ----
        scrap_note = ""
        if "scrap" in mp and "%" not in h["hn"][mp["scrap"]]:
            vals = [x["Hao hụt (%)"] for x in sheet_rows if x["Hao hụt (%)"] is not None and not x["_scrap_pct_fmt"]]
            if vals and all(0 < v <= 1 for v in vals):
                for x in sheet_rows:
                    f = x["Hao hụt (%)"]
                    if f is None or x["_scrap_pct_fmt"]:
                        continue
                    # Có cột KL nguyên liệu: hệ số đã nhân sẵn trong công thức KL -> không cộng thêm hao hụt
                    x["Hao hụt (%)"] = 0.0 if "weight" in mp else round((1 / f - 1) * 100, 6)
                scrap_note = " • Tỉ lệ hao hụt dạng hệ số" + (" (đã gồm trong KL nguyên liệu)" if "weight" in mp else "")
        # ---- Xác định cấp (Cấp) theo cột MỤC hoặc khoảng lùi đầu dòng ----
        has_dotted = any(re.fullmatch(DOTTED_RE, x["MỤC"]) and re.search(r"\d[.\-]\d", x["MỤC"]) for x in sheet_rows)
        ind_levels = sorted({x["_indent"] for x in sheet_rows})
        last_num_depth = 0
        for x in sheet_rows:
            m = x["MỤC"]
            if has_dotted:
                if re.fullmatch(DOTTED_RE, m):
                    d = len([p for p in re.split(r"[.\-]", m) if p])
                    last_num_depth = d
                elif m:
                    d = 0  # La Mã / chữ cái: nhóm tiêu đề
                else:
                    d = last_num_depth + 1
            else:
                if m and not re.fullmatch(DOTTED_RE, m):
                    d = 0
                else:
                    d = 1 + (ind_levels.index(x["_indent"]) if len(ind_levels) > 1 else 0)
            x["Cấp"] = float(d)
        all_rows.extend(sheet_rows)
        infos.append({"sheet": g["name"], "status": f"OK – {len(sheet_rows)} dòng, tiêu đề dòng {h['row'] + 1}"
                      + (" (2 tầng)" if h["two"] else "") + (f" • MỤC đa cột ({len(muc_cols)} cấp)" if len(muc_cols) > 1 else "")
                      + scrap_note,
                      "mapping": {k: h["hd"][v] for k, v in mp.items()},
                      "stages": [stage_names[k] for _, k in stage_cols]})
    if not all_rows:
        raise ValueError("Không đọc được dòng nào từ Technical BOM (cần cột MỤC/Mã chi tiết/Mô tả/Số lượng).")
    stage_keys = list(stage_names.keys())
    records = []
    for x in all_rows:
        rec = {k: x.get(k) for k in BOM_BASE_COLS}
        for k in stage_keys:
            rec[STAGE_PREFIX + stage_names[k]] = float(x["_stages"].get(k, 0.0))
        records.append(rec)
    df = pd.DataFrame(records)
    for c in NUM_COLS:
        df[c] = pd.to_numeric(df[c], errors="coerce").astype("float64")
    stage_cols = [STAGE_PREFIX + stage_names[k] for k in stage_keys]
    return df, stage_cols, sheet_ctx, infos, warns


# =============================================================================
# 6. ĐỊNH DANH NGUYÊN VẬT LIỆU THÔ (MỤC 2.3)
# =============================================================================


NON_METAL_KEYWORDS = ["vermiculite", "ceramic", "kính", "kinh", "gỗ", "go", "giấy", "giay", "nilon", "sợi", "soi", "keo", "silicagel"]
_NON_METAL_RE = re.compile(r"\b(?:" + "|".join(sorted({norm(k) for k in NON_METAL_KEYWORDS})) + r")\b")


def is_non_metal(desc):
    """Mô tả chứa từ khóa vật liệu phi kim loại (vermiculite, ceramic, kính...) – không phải phôi thép."""
    return bool(_NON_METAL_RE.search(norm(desc or "")))


def std_material_name(desc, material, t, w, l, w2=None):
    if is_non_metal(desc):
        return to_str(desc).strip()  # giữ nguyên tên gốc, không ép "Thép ..."
    d = norm(f"{desc} {material}")
    if not d:
        return None

    def from_desc(pat):
        m = re.search(pat, d)
        return to_float(m.group(1), None) if m else None

    if re.search(r"\bhop\b|\bbox\b|\brhs\b|\bshs\b|hop vuong|hop chu nhat", d):
        if w and w2 and t:  # BOM có đủ 2 cạnh tiết diện A x B -> quy cách thép hộp A x B x T
            return f"Thép hộp {fmt_dim(w)}x{fmt_dim(w2)}x{fmt_dim(t)}"
        if l and w and t:
            return f"Thép hộp {fmt_dim(l)}x{fmt_dim(w)}x{fmt_dim(t)}"
        m = re.search(r"(\d+(?:[.,]\d+)?)\s*x\s*(\d+(?:[.,]\d+)?)\s*x\s*(\d+(?:[.,]\d+)?)", d)
        if m:
            a, b, c = (to_float(m.group(i)) for i in (1, 2, 3))
            return f"Thép hộp {fmt_dim(a)}x{fmt_dim(b)}x{fmt_dim(c)}"
        return None
    if re.search(r"\bong\b|\bpipe\b|\btube\b", d):
        ww = w or from_desc(r"(?:phi|ø|o|d)\s*(\d+(?:[.,]\d+)?)")
        if ww and t:
            return f"Thép ống phi {fmt_dim(ww)}x{fmt_dim(t)}" + (f"x{fmt_dim(l)}" if l else "")
        return None
    if re.search(r"tron dac|\blap\b|round bar|\btron\b|\btruc\b", d):
        ww = w or from_desc(r"(?:phi|ø|d)\s*(\d+(?:[.,]\d+)?)")
        if ww:
            return f"Thép tròn đặc phi {fmt_dim(ww)}"
        return None
    if re.search(r"\btam\b|\bplate\b|\bsheet\b|\btole\b|\bton\b|\bpl\b|\bla thep\b", d):
        tt = t or from_desc(r"(?:\bt\s*=?\s*|day\s*|\bpl\s*)(\d+(?:[.,]\d+)?)")
        if tt:
            return f"Thép tấm {fmt_plate(tt)}"
        return None
    return None


# =============================================================================
# 7. BỘ MÁY RÃ BOM ĐA CẤP & MRP ĐỆ QUY (MỤC 2 & 4)
# =============================================================================

TYPE_LABEL = {"FG": "Thành phẩm (FG)", "SG": "Bán thành phẩm (SG)", "BTP": "BTP công đoạn (WIP)", "RM": "Nguyên vật liệu (RM)"}
TYPE_RANK = {"RM": 0, "SG": 1, "BTP": 1, "FG": 2}
# ĐVT đếm được / theo chiều dài của vật tư mua ngoài (đã norm)
COTS_UOMS = {"cai", "bo", "pcs", "pc", "set", "m", "met", "chiec", "con", "cap", "doi", "bich", "cuon", "lit", "chai"}


def _val(rec, k, default=None):
    v = rec.get(k)
    return default if is_blank(v) else v


_NUM_RE = re.compile(r"\d+(?:[.,]\d+)?")
_SHAPE_WORDS = ("tam", "hop", "ong", "tron")


def _nums(s):
    return [float(x.replace(",", ".")) for x in _NUM_RE.findall(s)]


def _same_nums(a, b):
    return len(a) == len(b) and all(abs(x - y) < 1e-6 for x, y in zip(a, b))


def safe_fuzzy_match(bom_name, po_name):
    """Khớp mờ an toàn tên NVL BOM ↔ tên PO: chỉ True khi chắc chắn cùng chủng loại và cùng kích thước số."""
    b, p = norm(bom_name or ""), norm(po_name or "")
    if not b or not p:
        return False
    has = lambda s, w: re.search(rf"\b{w}\b", s) is not None
    bn, pn = _nums(b), _nums(p)
    if not bn or not pn:
        return False
    if has(b, "thep") and has(b, "tam"):
        kind_ok = has(p, "tam") and bn[0] == pn[0]
        shapes = {"tam"}
    elif has(b, "thep") and has(b, "hop"):
        kind_ok = has(p, "hop") and _same_nums(sorted(bn), sorted(pn))
        shapes = {"hop"}
    elif (has(b, "thep") and has(b, "tron")) or has(b, "phi"):
        kind_ok = (has(p, "tron") or has(p, "phi")) and _same_nums(bn, pn)
        shapes = {w for w in _SHAPE_WORDS if has(b, w)}
    else:
        return False
    if not kind_ok:
        return False
    # PO mang thêm chủng loại khác (vd BOM tròn đặc – PO ống) → không chắc chắn → loại
    return all(w in shapes for w in _SHAPE_WORDS if has(p, w))


def run_engine(bom_df, stage_cols, sheet_ctx, so_df, po_df, P):
    W = []
    so_lines = so_df.to_dict("records")
    so_norm = {}
    for l in so_lines:
        so_norm.setdefault(ncode(l["Mã Sản Phẩm ERP"]), l["Mã Sản Phẩm ERP"])
    for l in so_lines:
        kh = ncode(l["Mã Sản Phẩm KH"])
        if kh and kh not in so_norm:
            so_norm[kh] = l["Mã Sản Phẩm ERP"]
    so_name = {l["Mã Sản Phẩm ERP"]: l["Tên Thành Phẩm"] for l in so_lines}

    so_suffix = [(ncode(l["Mã Sản Phẩm ERP"]), l["Mã Sản Phẩm ERP"]) for l in so_lines]

    def match_fg(v):
        """Khớp mã BOM -> mã FG trong SO: trùng Mã ERP, trùng Mã KH, hoặc Mã ERP có dạng
        '<tiền tố> - <mã BOM>' (VD: 'FG - SVC - SAV01829' <-> 'SAV01829')."""
        c = ncode(v)
        if not c:
            return None
        if c in so_norm:
            return so_norm[c]
        if len(c) >= 4:
            for k, erp in so_suffix:
                if k.endswith("-" + c):
                    return erp
        return None

    # ---- danh mục PO (kịch bản 3 tệp) ----
    po_map, po_by_name = {}, {}
    if po_df is not None and len(po_df):
        agg = {}
        for p in po_df.to_dict("records"):
            k = ncode(p["Mã Vật Tư ERP"]) or "NAME:" + norm(p["Tên Hàng Hóa Phụ Kiện Mua"])
            a = agg.setdefault(k, {"code": p["Mã Vật Tư ERP"], "name": p["Tên Hàng Hóa Phụ Kiện Mua"], "qty": 0.0,
                                   "amt": 0.0, "prices": [], "po": [], "sup": [], "date": p["Ngày Hẹn Giao"]})
            q, pr = p["Số Lượng Mua Thực Tế"], p["Đơn Giá Mua (VND)"]
            a["qty"] += q
            a["amt"] += q * pr
            a["prices"].append(pr)
            for fld, key in (("Số PO Mua", "po"), ("Nhà Cung Cấp", "sup")):
                if p[fld] and p[fld] not in a[key]:
                    a[key].append(p[fld])
        for k, a in agg.items():
            price = a["amt"] / a["qty"] if a["qty"] > 0 else (sum(a["prices"]) / len(a["prices"]) if a["prices"] else 0)
            rec = {"code": a["code"], "price": price, "po": ", ".join(a["po"]), "sup": ", ".join(a["sup"]),
                   "date": a["date"], "name": a["name"]}
            if not k.startswith("NAME:"):
                po_map[k] = rec
            if a["name"]:
                po_by_name.setdefault(norm(a["name"]), rec)

    prefix = ncode(P["rm_prefix"])
    default_scrap = P["default_scrap"]
    nodes, edges, expanded = {}, {}, {}
    stage_order = [c[len(STAGE_PREFIX):] for c in stage_cols]
    used_stages = []

    def reg(code, typ, name="", uom="", std="", stage="", price=0.0, seq=0):
        n = nodes.get(code)
        if n is None:
            n = nodes[code] = {"code": code, "type": typ, "name": name, "uom": uom, "std": std, "stage": stage,
                               "seq": seq, "price_bom": 0.0}
        else:
            if TYPE_RANK.get(typ, 0) > TYPE_RANK.get(n["type"], 0):
                n["type"] = typ
            for k, v in (("name", name), ("uom", uom), ("std", std), ("stage", stage)):
                if v and not n[k]:
                    n[k] = v
            if uom == "Kg" and n["uom"] != "Kg":
                W.append(f"NVL {code}: ĐVT không thống nhất ({n['uom']} / Kg) – chốt Kg.")
                n["uom"] = "Kg"
        if price and price > 0 and not n["price_bom"]:
            n["price_bom"] = price
        return n

    def add_edge(p, c, rate, scrap, occ):
        if expanded.setdefault(p, occ) != occ:
            return False
        e = edges.setdefault((p, c), {"rate": 0.0, "eff": 0.0})
        e["rate"] += rate
        e["eff"] += rate * (1 + scrap / 100.0)
        return True

    def rm_code_for(std, desc, material):
        if std and norm(std) in po_by_name and po_by_name[norm(std)]["code"]:
            return po_by_name[norm(std)]["code"]
        base = std or f"{material} {desc}".strip()
        return "RM-" + slug(base, "-")

    df = bom_df.copy()
    df["Sheet"] = df["Sheet"].replace("", None).ffill().fillna("Editor")
    records = df.to_dict("records")
    by_sheet = {}
    for rec in records:
        by_sheet.setdefault(rec["Sheet"], []).append(rec)

    trees = {}
    for sh, rows in by_sheet.items():
        parent, children, stack = [None] * len(rows), [[] for _ in rows], []
        depth = [int(to_float(_val(r, "Cấp"), 1)) for r in rows]
        for i in range(len(rows)):
            while stack and depth[stack[-1]] >= depth[i]:
                stack.pop()
            parent[i] = stack[-1] if stack else None
            if parent[i] is not None:
                children[parent[i]].append(i)
            stack.append(i)
        trees[sh] = (rows, children, [i for i in range(len(rows)) if parent[i] is None])

    def sheet_fg_of(sh):
        ctx = ncode(sheet_ctx.get(sh, sh))
        best = None
        for k, v in so_norm.items():
            if len(k) >= 3 and k in ctx and (best is None or len(k) > len(best[0])):
                best = (k, v)
        return best[1] if best else None

    def row_match(r):
        return match_fg(_val(r, "Mã chi tiết", "")) or None

    def subtree_has_match(rows, children, i):
        if row_match(rows[i]) or match_fg(_val(rows[i], "Mã TP", "")):
            return True
        return any(subtree_has_match(rows, children, c) for c in children[i])

    any_match = any(subtree_has_match(rows, ch, i) for rows, ch, roots in trees.values() for i in roots) or \
        any(sheet_fg_of(sh) for sh in trees)
    unique_fg = list(dict.fromkeys(l["Mã Sản Phẩm ERP"] for l in so_lines))
    single_fallback = unique_fg[0] if (len(unique_fg) == 1 and not any_match) else None
    if single_fallback:
        W.append(f"BOM không chứa mã FG nào khớp SO – gán toàn bộ BOM cho FG duy nhất '{single_fallback}'.")

    def row_stages(r):
        """Công đoạn được đánh dấu tại dòng, sắp theo số thứ tự (1, 2, 3...) rồi theo vị trí cột."""
        out = []
        for pos, c in enumerate(stage_cols):
            v = r.get(c)
            if truthy_mark(v):
                seq = to_float(v, 0)  # "x" -> 0 -> theo vị trí cột
                out.append((seq if seq > 0 else pos + 1, pos, c[len(STAGE_PREFIX):]))
        return [name for _, _, name in sorted(out)]

    anc_path = []  # mã (đã ncode) của mọi nút tổ tiên đang active trong nhánh hiện tại

    def process_inside(sh, rows, children, i, code, occ, scrap):
        if code in expanded and expanded[code] != occ:
            return  # nút dùng chung (DAG) – cấu trúc con đã được rã ở lần xuất hiện đầu tiên
        anc_path.append(ncode(code))
        try:
            _process_inside(sh, rows, children, i, code, occ, scrap)
        finally:
            anc_path.pop()

    def _process_inside(sh, rows, children, i, code, occ, scrap):
        r = rows[i]
        attach = code
        for stg in reversed(row_stages(r)):
            btp = f"{code}.BTP_{slug(stg)}"
            seq = stage_order.index(stg) + 1 if stg in stage_order else 0
            reg(btp, "BTP", name=f"{_val(r, 'Mô tả', code)} – sau {stg}", uom=_val(r, "ĐVT", "") or "Cái", stage=stg, seq=seq)
            if stg not in used_stages:
                used_stages.append(stg)
            add_edge(attach, btp, 1.0, 0.0, occ)
            attach = btp
        kg = to_float(_val(r, "KL NVL (kg)"), 0)
        desc, mat = _val(r, "Mô tả", ""), _val(r, "Vật liệu", "")
        std = std_material_name(desc, mat, to_float(_val(r, "Dày"), None), to_float(_val(r, "Rộng"), None),
                                to_float(_val(r, "Dài"), None), to_float(_val(r, "Rộng B"), None))
        rmc = _val(r, "Mã NVL", "")
        if rmc or kg > 0 or std:
            rmc = rmc or rm_code_for(std, desc, mat)
            # Namespace riêng cho NVL (như nhánh rm_leaf): tránh trùng mã với BTP/SG cha trước khi add_edge
            if not any(ncode(rmc).startswith(p) for p in (ncode(prefix), "RM-") if p):
                rmc = f"RM-{rmc}"
            uom = "Kg" if kg > 0 else (_val(r, "ĐVT", "") or "Cái")
            reg(rmc, "RM", name=std or f"{desc} {mat}".strip(), uom=uom, std=std or "", price=to_float(_val(r, "Đơn giá"), 0))
            add_edge(attach, rmc, kg if kg > 0 else 1.0, scrap, occ)
        for c in children[i]:
            process_row(sh, rows, children, c, attach, occ)

    def process_row(sh, rows, children, i, parent_code, parent_occ):
        r = rows[i]
        qty = to_float(_val(r, "SL/Cha"), 0)
        if qty <= 0:
            qty = 1.0
            if not _val(r, "SL/Cha"):
                W.append(f"[{sh}] Dòng {_val(r, 'Dòng', '?')} ({_val(r, 'Mã chi tiết', '') or _val(r, 'Mô tả', '')}): thiếu SL – mặc định 1.")
        scrap_v = _val(r, "Hao hụt (%)")
        scrap = to_float(scrap_v, default_scrap) if scrap_v is not None else default_scrap
        code = _val(r, "Mã chi tiết", "")
        if code and ncode(code) in anc_path:
            code = f"{code}_P"  # trùng bất kỳ cha/ông/cụ nào → đổi sang dạng Part để không bị ngắt "vòng lặp"
        kg = to_float(_val(r, "KL NVL (kg)"), 0)
        desc, mat = _val(r, "Mô tả", ""), _val(r, "Vật liệu", "")
        std = std_material_name(desc, mat, to_float(_val(r, "Dày"), None), to_float(_val(r, "Rộng"), None),
                                to_float(_val(r, "Dài"), None), to_float(_val(r, "Rộng B"), None))
        rm_col = _val(r, "Mã NVL", "")
        fabricated = bool(rm_col) or kg > 0 or bool(std)
        stages = row_stages(r)
        is_leaf = not children[i]
        code_is_rm = bool(code) and bool(prefix) and ncode(code).startswith(prefix)
        # Vật tư mua ngoài (COTS): ĐVT đếm được, không KL, không hình dạng phôi → RM thẳng, không sinh BTP
        # (kể cả khi dòng có gắn công đoạn; công đoạn vẫn được ghi nhận để tính TK 622/627)
        cots = (kg <= 0 and (not std or is_non_metal(desc)) and norm(_val(r, "ĐVT", "")) in COTS_UOMS)
        rm_leaf = is_leaf and (code_is_rm or cots or (not stages and (not code or not fabricated)))
        if rm_leaf:
            if not code and not desc and not rm_col:
                return
            for stg in (stages if cots else []):
                if stg not in used_stages:
                    used_stages.append(stg)
            rmc = code if code_is_rm else (rm_col or code or rm_code_for(std, desc, mat))
            # Namespace riêng cho NVL lá: tránh trùng mã với BTP/SG cha (vd GZ19726 → GZ19726) làm engine ngắt "vòng lặp"
            if not any(ncode(rmc).startswith(p) for p in (ncode(prefix), "RM-") if p):
                rmc = f"RM-{rmc}"
            uom = "Kg" if kg > 0 else (_val(r, "ĐVT", "") or "Cái")
            reg(rmc, "RM", name=std or desc or rmc, uom=uom, std=std or "", price=to_float(_val(r, "Đơn giá"), 0))
            add_edge(parent_code, rmc, qty * (kg if kg > 0 else 1.0), scrap, parent_occ)
            return
        code = code or f"{parent_code}-{_val(r, 'MỤC', '') or i + 1}"
        reg(code, "SG", name=desc or code, uom=_val(r, "ĐVT", "") or "Cái")
        add_edge(parent_code, code, qty, 0.0 if fabricated else scrap, parent_occ)
        process_inside(sh, rows, children, i, code, (sh, i), scrap)

    def resolve(rows, children, ids):
        out = []
        for i in ids:
            r = rows[i]
            fg = row_match(r)
            if fg:
                out.append((fg, i, True))
                continue
            fgc = match_fg(_val(r, "Mã TP", ""))
            if fgc:
                out.append((fgc, i, False))
                continue
            if children[i] and any(subtree_has_match(rows, children, c) for c in children[i]):
                out.extend(resolve(rows, children, children[i]))
                continue
            out.append((None, i, False))
        return out

    for sh, (rows, children, roots) in trees.items():
        sfg = sheet_fg_of(sh)
        for fg, i, is_fg_row in resolve(rows, children, roots):
            r = rows[i]
            if fg is None:
                if sfg or single_fallback:
                    fg = sfg or single_fallback
                elif children[i]:
                    fg = _val(r, "Mã chi tiết", "") or f"{sh}-{_val(r, 'MỤC', i + 1)}"
                    is_fg_row = True
                    W.append(f"[{sh}] Cụm gốc '{fg}' không có trong SO – nhu cầu = 0.")
                else:
                    fg = f"NO-FG-{slug(sh, '-')}"
                    W.append(f"[{sh}] Dòng {_val(r, 'Dòng', '?')} không xác định được Thành phẩm – gán vào '{fg}'.")
            reg(fg, "FG", name=so_name.get(fg, "") or _val(r, "Mô tả", ""), uom="Cái")
            if is_fg_row:
                scrap_v = _val(r, "Hao hụt (%)")
                process_inside(sh, rows, children, i, fg, (sh, i), to_float(scrap_v, default_scrap) if scrap_v is not None else default_scrap)
            else:
                anc_path.append(ncode(fg))
                try:
                    process_row(sh, rows, children, i, fg, ("fgctx", fg))
                finally:
                    anc_path.pop()

    # ---- chuẩn hóa loại nút ----
    ch = defaultdict(list)
    for (p, c), e in edges.items():
        ch[p].append((c, e["eff"]))
    for code, n in nodes.items():
        if n["type"] == "RM" and ch.get(code):
            n["type"] = "SG"
        if n["type"] in ("SG", "BTP") and not ch.get(code):
            W.append(f"Nút {code} ({n['name']}) không có NVL/chi tiết con – không phát sinh nhu cầu vật tư.")

    # ---- R_cum đệ quy theo mọi đường đi (DAG) ----
    memo, cycle_warned = {}, set()

    def cum(n, path):
        if n in memo:
            return memo[n]
        res = {n: 1.0}
        for c, r in ch.get(n, []):
            if c in path:
                if (n, c) not in cycle_warned:
                    cycle_warned.add((n, c))
                    W.append(f"Phát hiện vòng lặp BOM {n} → {c} – đã ngắt.")
                continue
            sub = cum(c, path | {c})
            for k, v in sub.items():
                res[k] = res.get(k, 0.0) + r * v
        memo[n] = res
        return res

    fg_cum = {}
    for l in so_lines:
        fg = l["Mã Sản Phẩm ERP"]
        if fg not in nodes:
            W.append(f"FG '{fg}' trong SO không có BOM – không phát sinh nhu cầu NVL.")
            nodes[fg] = {"code": fg, "type": "FG", "name": l["Tên Thành Phẩm"], "uom": l["ĐVT"], "std": "",
                         "stage": "", "seq": 0, "price_bom": 0.0}
        fg_cum[fg] = cum(fg, frozenset([fg]))

    rm_codes = [c for c, n in nodes.items() if n["type"] == "RM"]
    rm_used = [c for c in rm_codes if any(fg_cum[l["Mã Sản Phẩm ERP"]].get(c, 0) > 0 for l in so_lines)]
    R = [[fg_cum[l["Mã Sản Phẩm ERP"]].get(c, 0.0) for l in so_lines] for c in rm_used]

    # ---- Đơn giá: PO > BOM/Live Editor > Sidebar ----
    rm_rows = []
    for i, c in enumerate(rm_used):
        n = nodes[c]
        po = po_map.get(ncode(c)) or po_by_name.get(norm(n["name"])) or (po_by_name.get(norm(n["std"])) if n["std"] else None)
        if po is None and po_by_name:
            # Dự phòng: khớp mờ an toàn (chủng loại + kích thước số). Mơ hồ (nhiều giá khác nhau) → bỏ qua → 0đ.
            hits = [rec for rec in po_by_name.values()
                    if rec["price"] > 0 and (safe_fuzzy_match(n["std"] or n["name"], rec["name"])
                                             or safe_fuzzy_match(n["name"], rec["name"]))]
            if hits and len({round(h["price"], 2) for h in hits}) == 1:
                po = hits[0]
                W.append(f"Giá NVL {c} ({n['name']}) lấy theo khớp mờ an toàn với dòng PO '{po['name']}'.")
        # Thứ tự giá: (1) PO khớp > (2) BOM / Live Editor > (3) 0đ > (4) giá tạm Sidebar (chỉ khi người dùng bật).
        if po and po["price"] > 0:
            price, src = po["price"], f"PO: {po['po'] or 'không số'}"
        elif n["price_bom"] > 0:
            price, src = n["price_bom"], "BOM / Live Editor"
        else:
            price, src = 0.0, ("Chưa có giá PO (0đ)" if po_df is not None else "Chưa nạp PO (0đ)")
        if price <= 0 and P.get("apply_sidebar"):
            price = P["kg_price"] if n["uom"] == "Kg" else P["unit_price"]
            if price > 0:
                src = "Giá tạm Sidebar"
        if price <= 0:
            W.append(f"NVL {c} ({n['name']}, ĐVT {n['uom']}) có đơn giá 0đ – {src}.")
        q = sum(l["Số Lượng"] * R[i][j] for j, l in enumerate(so_lines))
        rm_rows.append({"Mã NVL": c, "Tên NVL": n["name"], "ĐVT": n["uom"], "Nhu cầu MRP": q, "Đơn giá (VND)": price,
                        "Thành tiền (VND)": q * price, "Nguồn giá": src,
                        "Số PO": po["po"] if po else "", "Nhà cung cấp": po["sup"] if po else "",
                        "Ngày hẹn giao": po["date"] if po else ""})
    mrp_df = pd.DataFrame(rm_rows, columns=["Mã NVL", "Tên NVL", "ĐVT", "Nhu cầu MRP", "Đơn giá (VND)", "Thành tiền (VND)",
                                            "Nguồn giá", "Số PO", "Nhà cung cấp", "Ngày hẹn giao"])

    # ---- WIP / DPP: thứ tự DFS từ FG ----
    wip = []
    for j, l in enumerate(so_lines):
        fg = l["Mã Sản Phẩm ERP"]
        seen, order = set(), []

        def go(n, d):
            for c, _ in ch.get(n, []):
                if c in seen or nodes[c]["type"] == "RM":
                    continue
                seen.add(c)
                order.append((c, d + 1))
                go(c, d + 1)

        go(fg, 0)
        maxd = max((d for _, d in order), default=0)
        for c, d in order:
            rate = fg_cum[fg].get(c, 0)
            if rate > 0:
                n = nodes[c]
                wip.append({"j": j, "code": c, "name": n["name"], "type": n["type"], "stage": n["stage"],
                            "rate": rate, "offset": maxd - d})

    return {"nodes": nodes, "edges": edges, "children": ch, "fg_cum": fg_cum, "so_lines": so_lines,
            "rm_used": rm_used, "R": R, "mrp_df": mrp_df, "wip": wip, "warnings": W,
            "used_stages": [s for s in stage_order if s in used_stages]}


def compute_financials(res, so_df, stage_rates, P):
    so_lines = res["so_lines"]
    mrp = res["mrp_df"]
    er = P["er"]
    total_qty = float(so_df["Số Lượng"].sum())
    labor_sum = sum(r["labor"] for r in stage_rates)
    oh_sum = sum(r["oh"] for r in stage_rates)
    prices = list(mrp["Đơn giá (VND)"]) if len(mrp) else []
    sku = []
    for j, l in enumerate(so_lines):
        mat = sum(res["R"][i][j] * prices[i] for i in range(len(prices)))
        unit = mat + labor_sum + oh_sum
        sell = l["Giá Bán FOB ($USD)"] * er
        sku.append({"Mã TP": l["Mã Sản Phẩm ERP"], "Tên TP": l["Tên Thành Phẩm"], "Số lượng": l["Số Lượng"],
                    "CP NVL/SP": mat, "CP NC/SP": labor_sum, "CP SXC/SP": oh_sum, "Giá thành/SP": unit,
                    "Tổng giá thành": unit * l["Số Lượng"], "Giá bán VND/SP": sell, "Lãi gộp/SP": sell - unit,
                    "Biên LN gộp": (sell - unit) / sell if sell else 0.0})
    tk621 = float(mrp["Thành tiền (VND)"].sum()) if len(mrp) else 0.0
    tk622 = total_qty * labor_sum
    tk627 = total_qty * oh_sum
    revenue = float(sum(l["Số Lượng"] * l["Giá Bán FOB ($USD)"] for l in so_lines)) * er
    cogs = tk621 + tk622 + tk627
    gp = revenue - cogs
    opex = revenue * P["sga"] / 100
    ebit = gp - opex
    tax = max(0.0, ebit * P["tax"] / 100)
    npat = max(0.0, ebit * (1 - P["tax"] / 100))
    pl = [("Doanh thu thuần (Revenue)", revenue), ("Chi phí NVL trực tiếp – TK 621", tk621),
          ("Chi phí nhân công trực tiếp – TK 622", tk622), ("Chi phí sản xuất chung – TK 627", tk627),
          ("Tổng giá vốn hàng bán (COGS)", cogs), ("Lợi nhuận gộp (Gross Profit)", gp),
          (f"Chi phí QL & BH (OPEX = {P['sga']:g}% DT)", opex), ("Lợi nhuận trước thuế & lãi vay (EBIT)", ebit),
          (f"Thuế TNDN ({P['tax']:g}%)", tax), ("Lợi nhuận thuần (NPAT)", npat)]
    return {"sku": pd.DataFrame(sku), "tk621": tk621, "tk622": tk622, "tk627": tk627, "revenue": revenue,
            "cogs": cogs, "gp": gp, "opex": opex, "ebit": ebit, "tax": tax, "npat": npat, "total_qty": total_qty,
            "pl": pd.DataFrame(pl, columns=["Chỉ tiêu", "Giá trị (VND)"]), "labor_sum": labor_sum, "oh_sum": oh_sum}


# =============================================================================
# 8. BỘ KẾT XUẤT EXCEL MASTER 25 SHEET – 100% CÔNG THỨC LIÊN KẾT SỐNG
# =============================================================================

SH = {
    1: "01. Thu muc & Thuat ngu", 2: "02. Bao gia & Don hang", 3: "03. KH Giao hang & Lenh SX",
    4: "04. Hoach dinh MRP", 5: "05. Lich ra chi tiet DPP", 6: "06. De nghi Mua hang PR",
    7: "07. So sanh Gia NCC", 8: "08. Don dat mua hang PO", 9: "09. Kiem tra NVL IQC",
    10: "10. Nhap kho NVL GRN", 11: "11. Xuat kho NVL OUT", 12: "12. Theo doi WIP",
    13: "13. Kiem tra TP FQC", 14: "14. Nhap kho TP FGRN", 15: "15. Lenh giao hang DO",
    16: "16. Phieu xuat kho GDN", 17: "17. Cong no phai thu AR", 18: "18. Danh sach Pallet ID",
    19: "19. Pha he Truy vet Nguoc", 20: "20. Tong hop Gia thanh", 21: "21. TK 621 CP NVL TT",
    22: "22. TK 622 CP NC TT", 23: "23. TK 627 CP SXC", 24: "24. Bao cao P&L (KQKD)", 25: "25. BOM ERP Da cap",
}
SH_DESC = {
    1: "Danh mục hệ thống & Thuật ngữ ERP", 2: "Bảng SO, ô B6 chứa tỷ giá, Thành tiền =G*H*$B$6",
    3: "Lệnh sản xuất (WO) kế thừa từ Sheet 02", 4: "Bảng rã MRP: Nhu cầu = SL SO × R_cum",
    5: "Lịch phân rã công đoạn dở dang WIP", 6: "Đề nghị mua hàng PR kế thừa Sheet 04",
    7: "So sánh báo giá nhà cung cấp", 8: "Đơn đặt mua hàng PO, đơn giá tra cứu PO/BOM",
    9: "Phiếu kiểm tra chất lượng NVL đầu vào (IQC)", 10: "Phiếu nhập kho NVL (GRN) + Lot NVL",
    11: "Phiếu xuất kho NVL cho sản xuất (OUT)", 12: "Theo dõi sản phẩm dở dang (WIP)",
    13: "Kiểm tra chất lượng thành phẩm (FQC)", 14: "Nhập kho thành phẩm (FGRN)",
    15: "Lệnh giao hàng (DO)", 16: "Phiếu xuất kho giao hàng (GDN)", 17: "Công nợ phải thu (AR) / Invoice",
    18: "Danh sách Pallet ID", 19: "Ma trận truy vết ngược 360° từ Invoice về Lot NVL",
    20: "Bảng tính giá thành sản xuất đích danh SKU", 21: "Hạch toán chi tiết TK 621",
    22: "Hạch toán chi tiết TK 622", 23: "Hạch toán chi tiết TK 627",
    24: "Báo cáo kết quả hoạt động kinh doanh dự án", 25: "BOM ERP đa cấp (FG → SG → BTP → RM)",
}
GLOSSARY = [
    ("FG", "Finished Goods – Thành phẩm (mã khóa SO)"), ("SG", "Semi-finished Goods – Bán thành phẩm / chi tiết"),
    ("BTP", "Bán thành phẩm công đoạn: {Mã_Cha}.BTP_{Tên_Công_Đoạn}"), ("RM", "Raw Material – Nguyên vật liệu thô"),
    ("BOM", "Bill of Materials – Định mức cấu trúc sản phẩm"), ("R_cum", "Tỷ lệ tiêu hao tích lũy nhân dồn theo mọi đường đi"),
    ("MRP", "Material Requirements Planning – Hoạch định nhu cầu vật tư"), ("WO", "Work Order – Lệnh sản xuất"),
    ("DPP", "Detailed Production Plan – Lịch rã chi tiết công đoạn"), ("PR", "Purchase Request – Đề nghị mua hàng"),
    ("PO", "Purchase Order – Đơn đặt mua hàng"), ("IQC", "Incoming Quality Control – Kiểm tra NVL đầu vào"),
    ("GRN", "Goods Received Note – Phiếu nhập kho NVL"), ("OUT", "Phiếu xuất kho NVL cho sản xuất"),
    ("WIP", "Work In Process – Sản phẩm dở dang"), ("FQC", "Final Quality Control – Kiểm tra thành phẩm"),
    ("FGRN", "Finished Goods Received Note – Nhập kho thành phẩm"), ("DO", "Delivery Order – Lệnh giao hàng"),
    ("GDN", "Goods Delivery Note – Phiếu xuất kho giao hàng"), ("AR", "Accounts Receivable – Công nợ phải thu"),
    ("TK 621", "Chi phí nguyên vật liệu trực tiếp"), ("TK 622", "Chi phí nhân công trực tiếp"),
    ("TK 627", "Chi phí sản xuất chung"), ("COGS", "Cost of Goods Sold – Giá vốn hàng bán"),
    ("EBIT", "Lợi nhuận trước thuế và lãi vay"), ("NPAT", "Net Profit After Tax – Lợi nhuận thuần"),
    ("FOB", "Free On Board – Giá bán giao lên tàu (USD)"),
]

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color="1F3864")
INP_FILL = PatternFill("solid", fgColor="FFF2CC")
NORM_FILL = PatternFill("solid", fgColor="E2EFDA")
TOT_FILL = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
F_INT, F_NUM, F_RATE, F_PCT, F_DATE = "#,##0", "#,##0.00", "#,##0.000000", "0.00%", "dd/mm/yyyy"


def q(k):
    return f"'{SH[k]}'!"


def L(c):
    return get_column_letter(c)


class XLWriter:
    def __init__(self, project, date_str):
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)
        self.project, self.date_str = project, date_str

    def put(self, ws, r, c, v, fmt=None, kind=None, bold=False, wrap=False):
        cell = safe_set_cell(ws, r, c, v)
        if cell is None:
            return None
        cell.border = BORDER
        if fmt:
            cell.number_format = fmt
        if kind == "in":
            cell.fill = INP_FILL
        elif kind == "norm":
            cell.fill = NORM_FILL
        elif kind == "tot":
            cell.fill = TOT_FILL
            bold = True
        if bold:
            cell.font = Font(bold=True)
        if wrap:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        return cell

    def sheet(self, k, headers, widths=None, subtitle=None):
        ws = self.wb.create_sheet(SH[k])
        ncols = max(len(headers), 6)
        last = L(ncols)
        ws.merge_cells(f"A1:{last}1")
        ws.merge_cells(f"A2:{last}2")
        ws.merge_cells(f"A3:{last}3")
        c = safe_set_cell(ws, "A1", value=f"METALIC – {SH[k][4:].upper()}")
        c.font = TITLE_FONT
        safe_set_cell(ws, "A2", value=subtitle or SH_DESC[k]).font = Font(italic=True, color="595959")
        safe_set_cell(ws, "A3", value=f"Dự án: {self.project}   |   Ngày lập: {self.date_str}   |   "
                                      f"Ô vàng = dữ liệu nhập • Ô xanh lá = định mức kỹ thuật • Còn lại = công thức sống")
        for i, h in enumerate(headers, 1):
            cell = self.put(ws, 8, i, h, wrap=True)
            cell.fill, cell.font = HDR_FILL, HDR_FONT
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.row_dimensions[8].height = 42
        for i, w in enumerate(widths or [], 1):
            ws.column_dimensions[L(i)].width = w
        ws.freeze_panes = "A9"
        return ws

    def save(self):
        self.wb.calculation.fullCalcOnLoad = True
        bio = io.BytesIO()
        self.wb.save(bio)
        return bio.getvalue()


DBL = Side(style="double", color="000000")
DBL_BORDER = Border(left=DBL, right=DBL, top=DBL, bottom=DBL)


def uom_total_label(u):
    return f"Tổng khối lượng ({u})" if str(u).strip().lower() == "kg" else f"Tổng chi tiết ({u})"


def total_row(X, ws, r, ncols, label_col, label, sum_cols, first=9):
    """Dòng TỔNG CỘNG chuẩn kế toán: công thức =SUM(đầu:cuối), chữ đậm, viền đôi."""
    for c in range(1, ncols + 1):
        X.put(ws, r, c, None, kind="tot").border = DBL_BORDER
    X.put(ws, r, label_col, label, kind="tot").border = DBL_BORDER
    for col, fmt in sum_cols.items():
        f = f"=SUM({L(col)}{first}:{L(col)}{r - 1})" if r - 1 >= first else "=0"
        X.put(ws, r, col, f, fmt, kind="tot").border = DBL_BORDER


def build_master_excel(res, fin, so_df, stage_rates, P):
    X = XLWriter(P["project"], P["start_date"].strftime("%d/%m/%Y"))
    so = res["so_lines"]
    nso = len(so)
    mrp = res["mrp_df"].to_dict("records")
    nrm = len(mrp)
    R = res["R"]
    st_rows = stage_rates
    nst = len(st_rows)
    tot2 = 9 + nso
    tot22 = 9 + nst  # dòng tổng của Sheet 22/23
    TOTAL_QTY = f"{q(2)}$G${tot2}"
    ER = f"{q(2)}$B$6"
    rm_last = 8 + nrm
    yymm = P["start_date"].strftime("%y%m")
    put = X.put

    # ---------------- 01. Thư mục & Thuật ngữ ----------------
    ws = X.sheet(1, ["STT", "Tên Sheet", "Chức năng & Công thức liên kết"], [8, 34, 70])
    for k in range(1, 26):
        r = 8 + k
        put(ws, r, 1, k)
        put(ws, r, 2, f'=HYPERLINK("#\'{SH[k]}\'!A1","{SH[k]}")')
        put(ws, r, 3, SH_DESC[k])
    r0 = 36
    for i, h in enumerate(["STT", "Thuật ngữ", "Giải nghĩa"], 1):
        c = put(ws, r0, i, h)
        c.fill, c.font = HDR_FILL, HDR_FONT
    for i, (t, d) in enumerate(GLOSSARY, 1):
        put(ws, r0 + i, 1, i)
        put(ws, r0 + i, 2, t, bold=True)
        put(ws, r0 + i, 3, d)
    safe_set_cell(ws, "A4", value="Tỷ giá USD/VND (liên kết '02'!B6):")
    put(ws, 4, 3, f"={ER}", F_INT)
    safe_set_cell(ws, "A5", value="Tổng sản lượng SO:")
    put(ws, 5, 3, f"={TOTAL_QTY}", F_INT)
    safe_set_cell(ws, "A6", value="Lợi nhuận thuần NPAT (VND):")
    put(ws, 6, 3, f"={q(24)}$D$18", F_INT)

    # ---------------- 02. Báo giá & Đơn hàng ----------------
    ws = X.sheet(2, ["STT", "Mã SP KH", "Mã SP ERP", "Tên thành phẩm", "Quy cách", "ĐVT", "Số lượng",
                     "Giá bán FOB (USD)", "Thành tiền (VND)", "Thành tiền (USD)"], [6, 16, 18, 34, 20, 8, 12, 14, 20, 16])
    safe_set_cell(ws, "A4", value="Khách hàng / Dự án:")
    safe_set_cell(ws, "B4", value=P["project"])
    safe_set_cell(ws, "A5", value="Ngày đơn hàng:")
    c = safe_set_cell(ws, "B5", value=P["start_date"])
    c.number_format = F_DATE
    safe_set_cell(ws, "A6", value="Tỷ giá USD/VND:")
    put(ws, 6, 2, P["er"], F_INT, kind="in")
    safe_set_cell(ws, "A7", value="Tổng SL SO:")
    put(ws, 7, 2, f"=G{tot2}", F_INT)
    for j, l in enumerate(so):
        r = 9 + j
        put(ws, r, 1, j + 1)
        put(ws, r, 2, l["Mã Sản Phẩm KH"])
        put(ws, r, 3, l["Mã Sản Phẩm ERP"])
        put(ws, r, 4, l["Tên Thành Phẩm"])
        put(ws, r, 5, l["Quy Cách"])
        put(ws, r, 6, l["ĐVT"])
        put(ws, r, 7, l["Số Lượng"], F_INT, kind="in")
        put(ws, r, 8, l["Giá Bán FOB ($USD)"], F_NUM, kind="in")
        put(ws, r, 9, f"=G{r}*H{r}*$B$6", F_INT)
        put(ws, r, 10, f"=G{r}*H{r}", F_NUM)
    total_row(X, ws, tot2, 10, 4, "TỔNG CỘNG", {7: F_INT, 9: F_INT, 10: F_NUM})
    REVENUE = f"{q(2)}$I${tot2}"

    # ---------------- 03. KH Giao hàng & Lệnh SX ----------------
    ws = X.sheet(3, ["STT", "Số lệnh SX (WO)", "Mã TP", "Tên TP", "SL đơn hàng", "% dự phòng SX", "SL lệnh SX",
                     "Ngày bắt đầu", "Lead time (ngày)", "Ngày hoàn thành", "Ngày giao (ETD)"],
                 [6, 18, 18, 34, 12, 12, 12, 13, 12, 14, 14])
    for j in range(nso):
        r = 9 + j
        put(ws, r, 1, j + 1)
        put(ws, r, 2, f"WO-{yymm}-{j + 1:03d}")
        put(ws, r, 3, f"={q(2)}C{r}")
        put(ws, r, 4, f"={q(2)}D{r}")
        put(ws, r, 5, f"={q(2)}G{r}", F_INT)
        put(ws, r, 6, 0, F_PCT, kind="in")
        put(ws, r, 7, f"=ROUNDUP(E{r}*(1+F{r}),0)", F_INT)
        put(ws, r, 8, P["start_date"], F_DATE, kind="in")
        put(ws, r, 9, P["lead_days"], F_INT, kind="in")
        put(ws, r, 10, f"=H{r}+I{r}", F_DATE)
        put(ws, r, 11, f"=J{r}+2", F_DATE)

    total_row(X, ws, 9 + nso, 11, 4, "TỔNG CỘNG", {5: F_INT, 7: F_INT})

    # ---------------- 04. Hoạch định MRP ----------------
    hdr = ["STT", "Mã NVL", "Tên NVL", "ĐVT", "Nhu cầu MRP", "Số FG sử dụng"] + \
          [f"R_cum ← {l['Mã Sản Phẩm ERP']}" for l in so]
    ws = X.sheet(4, hdr, [6, 24, 34, 8, 16, 10] + [16] * nso,
                 subtitle="Q_MRP(RM) = Σ Q_SO(FG) × R_cum(RM, FG) – liên kết sống SL từ Sheet 02")
    safe_set_cell(ws, "E6", value="Định mức tích lũy R_cum (ĐVT NVL / 1 TP) →")
    safe_set_cell(ws, "E7", value="SL SO liên kết Sheet 02 →")
    for j in range(nso):
        put(ws, 7, 7 + j, f"={q(2)}$G${9 + j}", F_INT)
    for i, m in enumerate(mrp):
        r = 9 + i
        put(ws, r, 1, i + 1)
        put(ws, r, 2, m["Mã NVL"])
        put(ws, r, 3, m["Tên NVL"])
        put(ws, r, 4, m["ĐVT"])
        terms = [f"{q(2)}$G${9 + j}*{L(7 + j)}{r}" for j in range(nso) if R[i][j] > 0]
        put(ws, r, 5, "=" + ("+".join(terms) if terms else "0"), F_NUM, bold=True)
        put(ws, r, 6, f'=COUNTIF({L(7)}{r}:{L(6 + nso)}{r},">0")', F_INT)
        for j in range(nso):
            put(ws, r, 7 + j, R[i][j], F_RATE, kind="norm")
    # Tổng riêng theo từng ĐVT (không cộng lẫn Kg với Cái): =SUMIF(cột ĐVT, "Kg", cột Nhu cầu MRP)
    for k, u in enumerate(dict.fromkeys(m["ĐVT"] for m in mrp)):
        r = 9 + nrm + k
        for c in range(1, 7):
            X.put(ws, r, c, None, kind="tot").border = DBL_BORDER
        X.put(ws, r, 3, uom_total_label(u), kind="tot").border = DBL_BORDER
        X.put(ws, r, 4, u, kind="tot").border = DBL_BORDER
        X.put(ws, r, 5, f'=SUMIF($D$9:$D${rm_last},D{r},E$9:E${rm_last})', F_NUM, kind="tot").border = DBL_BORDER

    # ---------------- 05. Lịch rã chi tiết DPP ----------------
    wip = res["wip"]
    ws = X.sheet(5, ["STT", "Số WO", "Mã TP", "Mã BTP / Chi tiết", "Tên", "Công đoạn", "Định mức tích lũy",
                     "SL kế hoạch", "Lùi lịch (ngày)", "Ngày bắt đầu", "Ngày kết thúc"],
                 [6, 16, 18, 34, 36, 14, 14, 14, 10, 13, 13])
    for k, w in enumerate(wip):
        r, rj = 9 + k, 9 + w["j"]
        put(ws, r, 1, k + 1)
        put(ws, r, 2, f"={q(3)}B{rj}")
        put(ws, r, 3, f"={q(2)}C{rj}")
        put(ws, r, 4, w["code"])
        put(ws, r, 5, w["name"])
        put(ws, r, 6, w["stage"] or ("Lắp ráp chi tiết" if w["type"] == "SG" else ""))
        put(ws, r, 7, w["rate"], F_RATE, kind="norm")
        put(ws, r, 8, f"={q(2)}$G${rj}*G{r}", F_NUM)
        put(ws, r, 9, w["offset"], F_INT, kind="in")
        put(ws, r, 10, f"={q(3)}$H${rj}+I{r}", F_DATE)
        put(ws, r, 11, f"=J{r}+1", F_DATE)
    wip_last = 8 + len(wip)
    total_row(X, ws, 9 + len(wip), 11, 5, "TỔNG CỘNG", {8: F_NUM})

    # ---------------- 06. Đề nghị mua hàng PR ----------------
    ws = X.sheet(6, ["STT", "Số PR", "Mã NVL", "Tên NVL", "ĐVT", "Nhu cầu MRP", "Tồn kho khả dụng",
                     "SL đề nghị mua", "Ngày cần hàng"], [6, 16, 24, 34, 8, 16, 14, 16, 14])
    for i in range(nrm):
        r = 9 + i
        put(ws, r, 1, i + 1)
        put(ws, r, 2, f"PR-{yymm}-{i + 1:03d}")
        put(ws, r, 3, f"={q(4)}B{r}")
        put(ws, r, 4, f"={q(4)}C{r}")
        put(ws, r, 5, f"={q(4)}D{r}")
        put(ws, r, 6, f"={q(4)}E{r}", F_NUM)
        put(ws, r, 7, 0, F_NUM, kind="in")
        put(ws, r, 8, f"=MAX(0,F{r}-G{r})", F_NUM)
        put(ws, r, 9, f"=MIN({q(3)}$H$9:$H${8 + nso})-3", F_DATE)
    total_row(X, ws, 9 + nrm, 9, 4, "TỔNG CỘNG", {6: F_NUM, 7: F_NUM, 8: F_NUM})

    # ---------------- 07. So sánh giá NCC ----------------
    ws = X.sheet(7, ["STT", "Mã NVL", "Tên NVL", "ĐVT", "SL đề nghị", "Giá NCC chính (PO/BOM)", "Giá NCC 2",
                     "Giá NCC 3", "Giá thấp nhất", "NCC đề xuất", "Thành tiền theo giá thấp nhất", "Tiết kiệm so với PO"],
                 [6, 24, 34, 8, 14, 16, 14, 14, 14, 22, 20, 18])
    for i in range(nrm):
        r = 9 + i
        put(ws, r, 1, i + 1)
        put(ws, r, 2, f"={q(6)}C{r}")
        put(ws, r, 3, f"={q(6)}D{r}")
        put(ws, r, 4, f"={q(6)}E{r}")
        put(ws, r, 5, f"={q(6)}H{r}", F_NUM)
        put(ws, r, 6, f"={q(8)}H{r}", F_INT)
        put(ws, r, 7, None, F_INT, kind="in")
        put(ws, r, 8, None, F_INT, kind="in")
        put(ws, r, 9, f"=MIN(F{r}:H{r})", F_INT)
        put(ws, r, 10, f'=IF(I{r}=F{r},IF({q(8)}C{r}="","NCC chính",{q(8)}C{r}),IF(I{r}=G{r},"NCC 2","NCC 3"))')
        put(ws, r, 11, f"=E{r}*I{r}", F_INT)
        put(ws, r, 12, f"=(F{r}-I{r})*E{r}", F_INT)
    total_row(X, ws, 9 + nrm, 12, 3, "TỔNG CỘNG", {5: F_NUM, 11: F_INT, 12: F_INT})

    # ---------------- 08. Đơn đặt mua hàng PO ----------------
    tot8 = 9 + nrm
    ws = X.sheet(8, ["STT", "Số PO", "Nhà cung cấp", "Mã NVL", "Tên NVL", "ĐVT", "SL đặt mua", "Đơn giá (VND)",
                     "Thành tiền (VND)", "Nguồn đơn giá", "Ngày hẹn giao"], [6, 18, 24, 24, 34, 8, 14, 14, 18, 20, 14])
    for i, m in enumerate(mrp):
        r = 9 + i
        put(ws, r, 1, i + 1)
        put(ws, r, 2, m["Số PO"] or f"PO-{yymm}-{i + 1:03d}", kind="in" if m["Số PO"] else None)
        put(ws, r, 3, m["Nhà cung cấp"], kind="in")
        put(ws, r, 4, f"={q(6)}C{r}")
        put(ws, r, 5, f"={q(6)}D{r}")
        put(ws, r, 6, f"={q(6)}E{r}")
        put(ws, r, 7, f"={q(6)}H{r}", F_NUM)
        put(ws, r, 8, m["Đơn giá (VND)"], F_INT, kind="in")
        put(ws, r, 9, f"=G{r}*H{r}", F_INT)
        put(ws, r, 10, m["Nguồn giá"])
        put(ws, r, 11, m["Ngày hẹn giao"] if m["Ngày hẹn giao"] else f"={q(6)}I{r}",
            None if m["Ngày hẹn giao"] else F_DATE)
    total_row(X, ws, tot8, 11, 5, "TỔNG GIÁ TRỊ PO", {7: F_NUM, 9: F_INT})

    # ---------------- 09. IQC ----------------
    ws = X.sheet(9, ["STT", "Số phiếu IQC", "Số PO", "Mã NVL", "Tên NVL", "SL nhận", "SL đạt", "SL không đạt",
                     "Tỷ lệ đạt", "Kết luận"], [6, 16, 18, 24, 34, 14, 14, 14, 10, 14])
    for i in range(nrm):
        r = 9 + i
        put(ws, r, 1, i + 1)
        put(ws, r, 2, f"IQC-{yymm}-{i + 1:03d}")
        put(ws, r, 3, f"={q(8)}B{r}")
        put(ws, r, 4, f"={q(8)}D{r}")
        put(ws, r, 5, f"={q(8)}E{r}")
        put(ws, r, 6, f"={q(8)}G{r}", F_NUM)
        put(ws, r, 7, f"=F{r}", F_NUM)
        put(ws, r, 8, f"=F{r}-G{r}", F_NUM)
        put(ws, r, 9, f"=IF(F{r}=0,0,G{r}/F{r})", F_PCT)
        put(ws, r, 10, f'=IF(H{r}=0,"ĐẠT","KHÔNG ĐẠT")')
    total_row(X, ws, 9 + nrm, 10, 5, "TỔNG CỘNG", {6: F_NUM, 7: F_NUM, 8: F_NUM})

    # ---------------- 10. GRN ----------------
    ws = X.sheet(10, ["STT", "Số phiếu GRN", "Mã NVL", "Tên NVL", "SL nhập kho", "Đơn giá", "Lot NVL", "Thành tiền"],
                 [6, 16, 24, 34, 14, 14, 22, 18])
    for i, m in enumerate(mrp):
        r = 9 + i
        put(ws, r, 1, i + 1)
        put(ws, r, 2, f"GRN-{yymm}-{i + 1:03d}")
        put(ws, r, 3, f"={q(9)}D{r}")
        put(ws, r, 4, f"={q(9)}E{r}")
        put(ws, r, 5, f"={q(9)}G{r}", F_NUM)
        put(ws, r, 6, f"={q(8)}H{r}", F_INT)
        put(ws, r, 7, f"LOT-{yymm}-{i + 1:04d}")
        put(ws, r, 8, f"=E{r}*F{r}", F_INT)
    total_row(X, ws, 9 + nrm, 8, 4, "TỔNG CỘNG", {5: F_NUM, 8: F_INT})

    # ---------------- 11. OUT ----------------
    ws = X.sheet(11, ["STT", "Số phiếu xuất", "Mã NVL", "Tên NVL", "SL xuất SX", "Đơn giá", "Thành tiền",
                      "Tồn kho sau xuất", "Lot NVL"], [6, 16, 24, 34, 14, 14, 18, 14, 22])
    for i in range(nrm):
        r = 9 + i
        put(ws, r, 1, i + 1)
        put(ws, r, 2, f"OUT-{yymm}-{i + 1:03d}")
        put(ws, r, 3, f"={q(4)}B{r}")
        put(ws, r, 4, f"={q(4)}C{r}")
        put(ws, r, 5, f"={q(4)}E{r}", F_NUM)
        put(ws, r, 6, f"={q(8)}H{r}", F_INT)
        put(ws, r, 7, f"=E{r}*F{r}", F_INT)
        put(ws, r, 8, f"={q(10)}E{r}-E{r}", F_NUM)
        put(ws, r, 9, f"={q(10)}G{r}")
    total_row(X, ws, 9 + nrm, 9, 4, "TỔNG CỘNG", {5: F_NUM, 7: F_INT})

    # ---------------- 12. WIP ----------------
    ws = X.sheet(12, ["STT", "Mã BTP / Chi tiết", "Tên", "Công đoạn", "SL kế hoạch", "SL hoàn thành", "SL dở dang",
                      "% hoàn thành"], [6, 34, 36, 14, 14, 14, 14, 12])
    uniq = list(dict.fromkeys(w["code"] for w in wip))
    nodes = res["nodes"]
    for k, code in enumerate(uniq):
        r = 9 + k
        put(ws, r, 1, k + 1)
        put(ws, r, 2, code)
        put(ws, r, 3, nodes[code]["name"])
        put(ws, r, 4, nodes[code]["stage"])
        put(ws, r, 5, f"=SUMIF({q(5)}$D$9:$D${wip_last},B{r},{q(5)}$H$9:$H${wip_last})", F_NUM)
        put(ws, r, 6, 0, F_NUM, kind="in")
        put(ws, r, 7, f"=E{r}-F{r}", F_NUM)
        put(ws, r, 8, f"=IF(E{r}=0,0,F{r}/E{r})", F_PCT)
    total_row(X, ws, 9 + len(uniq), 8, 3, "TỔNG CỘNG", {5: F_NUM, 6: F_NUM, 7: F_NUM})

    # ---------------- 13 → 18: chứng từ thành phẩm ----------------
    ws13 = X.sheet(13, ["STT", "Số phiếu FQC", "Số WO", "Mã TP", "SL kiểm tra", "SL đạt", "SL lỗi", "Tỷ lệ đạt"],
                   [6, 16, 16, 18, 14, 14, 14, 10])
    ws14 = X.sheet(14, ["STT", "Số phiếu FGRN", "Mã TP", "SL nhập kho TP", "Giá thành đơn vị", "Giá trị nhập kho"],
                   [6, 16, 18, 14, 16, 20])
    ws15 = X.sheet(15, ["STT", "Số DO", "Mã TP", "Tên TP", "SL giao", "Ngày giao"], [6, 16, 18, 34, 14, 14])
    ws16 = X.sheet(16, ["STT", "Số GDN", "Số DO", "Mã TP", "SL xuất giao", "Tồn TP sau xuất"], [6, 16, 16, 18, 14, 14])
    ws17 = X.sheet(17, ["STT", "Số Invoice", "Số GDN", "Mã TP", "SL", "Giá FOB (USD)", "Thành tiền (USD)",
                        "Quy đổi (VND)", "Hạn thanh toán", "Đã thu (VND)", "Còn phải thu (VND)"],
                   [6, 16, 16, 18, 12, 12, 16, 20, 14, 18, 20])
    ws18 = X.sheet(18, ["STT", "Mã TP", "SL giao", "SL / Pallet", "Số pallet", "Pallet ID đầu", "Pallet ID cuối"],
                   [6, 18, 12, 12, 10, 18, 18])
    for j in range(nso):
        r = 9 + j
        for w in (ws13, ws14, ws15, ws16, ws17, ws18):
            put(w, r, 1, j + 1)
        put(ws13, r, 2, f"FQC-{yymm}-{j + 1:03d}")
        put(ws13, r, 3, f"={q(3)}B{r}")
        put(ws13, r, 4, f"={q(3)}C{r}")
        put(ws13, r, 5, f"={q(3)}G{r}", F_INT)
        put(ws13, r, 6, f"=E{r}", F_INT)
        put(ws13, r, 7, f"=E{r}-F{r}", F_INT)
        put(ws13, r, 8, f"=IF(E{r}=0,0,F{r}/E{r})", F_PCT)
        put(ws14, r, 2, f"FGRN-{yymm}-{j + 1:03d}")
        put(ws14, r, 3, f"={q(13)}D{r}")
        put(ws14, r, 4, f"={q(13)}F{r}", F_INT)
        put(ws14, r, 5, f"={q(20)}H{r}", F_INT)
        put(ws14, r, 6, f"=D{r}*E{r}", F_INT)
        put(ws15, r, 2, f"DO-{yymm}-{j + 1:03d}")
        put(ws15, r, 3, f"={q(2)}C{r}")
        put(ws15, r, 4, f"={q(2)}D{r}")
        put(ws15, r, 5, f"={q(2)}G{r}", F_INT)
        put(ws15, r, 6, f"={q(3)}K{r}", F_DATE)
        put(ws16, r, 2, f"GDN-{yymm}-{j + 1:03d}")
        put(ws16, r, 3, f"={q(15)}B{r}")
        put(ws16, r, 4, f"={q(15)}C{r}")
        put(ws16, r, 5, f"=MIN({q(15)}E{r},{q(14)}D{r})", F_INT)
        put(ws16, r, 6, f"={q(14)}D{r}-E{r}", F_INT)
        put(ws17, r, 2, f"INV-{yymm}-{j + 1:03d}")
        put(ws17, r, 3, f"={q(16)}B{r}")
        put(ws17, r, 4, f"={q(16)}D{r}")
        put(ws17, r, 5, f"={q(16)}E{r}", F_INT)
        put(ws17, r, 6, f"={q(2)}H{r}", F_NUM)
        put(ws17, r, 7, f"=E{r}*F{r}", F_NUM)
        put(ws17, r, 8, f"=G{r}*{ER}", F_INT)
        put(ws17, r, 9, f"={q(15)}F{r}+30", F_DATE)
        put(ws17, r, 10, 0, F_INT, kind="in")
        put(ws17, r, 11, f"=H{r}-J{r}", F_INT)
        put(ws18, r, 2, f"={q(15)}C{r}")
        put(ws18, r, 3, f"={q(15)}E{r}", F_INT)
        put(ws18, r, 4, P["pallet_qty"], F_INT, kind="in")
        put(ws18, r, 5, f"=IF(D{r}=0,0,ROUNDUP(C{r}/D{r},0))", F_INT)
        put(ws18, r, 6, f'="PLT-"&TEXT(A{r},"000")&"-001"')
        put(ws18, r, 7, f'="PLT-"&TEXT(A{r},"000")&"-"&TEXT(MAX(1,E{r}),"000")')
    tot17 = 9 + nso
    total_row(X, ws13, tot17, 8, 4, "TỔNG CỘNG", {5: F_INT, 6: F_INT, 7: F_INT})
    total_row(X, ws14, tot17, 6, 3, "TỔNG CỘNG", {4: F_INT, 6: F_INT})
    total_row(X, ws15, tot17, 6, 4, "TỔNG CỘNG", {5: F_INT})
    total_row(X, ws16, tot17, 6, 4, "TỔNG CỘNG", {5: F_INT, 6: F_INT})
    total_row(X, ws17, tot17, 11, 4, "TỔNG CỘNG", {5: F_INT, 7: F_NUM, 8: F_INT, 10: F_INT, 11: F_INT})
    total_row(X, ws18, tot17, 7, 2, "TỔNG CỘNG", {3: F_INT, 5: F_INT})

    # ---------------- 19. Phả hệ truy vết ngược ----------------
    ws = X.sheet(19, ["STT", "Số Invoice", "Số DO", "Mã TP", "Số WO", "Pallet ID", "Mã NVL", "Tên NVL", "Lot NVL",
                      "Số PO", "Nhà cung cấp", "SL NVL tiêu hao"], [6, 16, 16, 18, 16, 26, 24, 30, 18, 18, 22, 16],
                 subtitle="Ma trận truy vết ngược 360°: Invoice → DO → TP → WO → Pallet → NVL → Lot → PO → NCC")
    k = 0
    for j in range(nso):
        rj = 9 + j
        for i in range(nrm):
            if R[i][j] <= 0:
                continue
            r, ri = 9 + k, 9 + i
            put(ws, r, 1, k + 1)
            put(ws, r, 2, f"={q(17)}B{rj}")
            put(ws, r, 3, f"={q(15)}B{rj}")
            put(ws, r, 4, f"={q(2)}C{rj}")
            put(ws, r, 5, f"={q(3)}B{rj}")
            put(ws, r, 6, f'={q(18)}F{rj}&" → "&{q(18)}G{rj}')
            put(ws, r, 7, f"={q(4)}B{ri}")
            put(ws, r, 8, f"={q(4)}C{ri}")
            put(ws, r, 9, f"={q(10)}G{ri}")
            put(ws, r, 10, f"={q(8)}B{ri}")
            put(ws, r, 11, f"={q(8)}C{ri}")
            put(ws, r, 12, f"={q(2)}$G${rj}*{q(4)}{L(7 + j)}{ri}", F_NUM)
            k += 1
    total_row(X, ws, 9 + k, 12, 8, "TỔNG CỘNG", {12: F_NUM})

    # ---------------- 20. Tổng hợp giá thành ----------------
    ws = X.sheet(20, ["STT", "Mã TP", "Tên TP", "Số lượng", "CP NVL / SP", "CP NC / SP", "CP SXC / SP",
                      "Giá thành đơn vị", "Tổng giá thành", "Giá bán VND / SP", "Lãi gộp / SP", "Biên LN gộp"],
                 [6, 18, 34, 12, 16, 14, 14, 16, 20, 16, 16, 12])
    for j in range(nso):
        r = 9 + j
        col = L(7 + j)
        put(ws, r, 1, j + 1)
        put(ws, r, 2, f"={q(2)}C{r}")
        put(ws, r, 3, f"={q(2)}D{r}")
        put(ws, r, 4, f"={q(2)}G{r}", F_INT)
        put(ws, r, 5, f"=SUMPRODUCT({q(4)}{col}$9:{col}${rm_last},{q(8)}$H$9:$H${rm_last})" if nrm else "=0", F_INT)
        put(ws, r, 6, f"={q(22)}$C${tot22}", F_INT)
        put(ws, r, 7, f"={q(23)}$C${tot22}", F_INT)
        put(ws, r, 8, f"=E{r}+F{r}+G{r}", F_INT)
        put(ws, r, 9, f"=D{r}*H{r}", F_INT)
        put(ws, r, 10, f"={q(2)}H{r}*{ER}", F_INT)
        put(ws, r, 11, f"=J{r}-H{r}", F_INT)
        put(ws, r, 12, f"=IF(J{r}=0,0,K{r}/J{r})", F_PCT)
    tot20 = 9 + nso
    total_row(X, ws, tot20, 12, 3, "TỔNG CỘNG", {4: F_INT, 9: F_INT})

    # ---------------- 21. TK 621 ----------------
    tot21 = 9 + nrm
    ws = X.sheet(21, ["STT", "Mã NVL", "Tên NVL", "ĐVT", "SL tiêu hao (MRP)", "Đơn giá", "Thành tiền (VND)",
                      "Nguồn đơn giá"], [6, 24, 34, 8, 16, 14, 20, 20],
                 subtitle="TK 621 = Σ [ Q_MRP(RM) × Đơn giá PO/BOM(RM) ]")
    for i in range(nrm):
        r = 9 + i
        put(ws, r, 1, i + 1)
        put(ws, r, 2, f"={q(4)}B{r}")
        put(ws, r, 3, f"={q(4)}C{r}")
        put(ws, r, 4, f"={q(4)}D{r}")
        put(ws, r, 5, f"={q(4)}E{r}", F_NUM)
        put(ws, r, 6, f"={q(8)}H{r}", F_INT)
        put(ws, r, 7, f"=E{r}*F{r}", F_INT)
        put(ws, r, 8, f"={q(8)}J{r}")
    total_row(X, ws, tot21, 8, 3, "TỔNG CỘNG TK 621", {5: F_NUM, 7: F_INT})

    # ---------------- 22 & 23. TK 622 / TK 627 ----------------
    for key, field, label in ((22, "labor", "Đơn giá nhân công / SP (VND)"), (23, "oh", "Đơn giá SXC / SP (VND)")):
        ws = X.sheet(key, ["STT", "Công đoạn", label, "Tổng SL SO", "Thành tiền (VND)"], [6, 26, 22, 14, 20],
                     subtitle=f"TK {'622' if key == 22 else '627'} = Σ_Stage [ Tổng SL SO × Đơn giá công đoạn ]")
        for s, row in enumerate(st_rows):
            r = 9 + s
            put(ws, r, 1, s + 1)
            put(ws, r, 2, row["stage"])
            put(ws, r, 3, row[field], F_INT, kind="in")
            put(ws, r, 4, f"={TOTAL_QTY}", F_INT)
            put(ws, r, 5, f"=C{r}*D{r}", F_INT)
        total_row(X, ws, tot22, 5, 2, "TỔNG CỘNG", {3: F_INT, 5: F_INT})

    # ---------------- 24. Báo cáo P&L ----------------
    ws = X.sheet(24, ["STT", "Chỉ tiêu", "Diễn giải công thức", "Giá trị (VND)", "% Doanh thu"], [6, 40, 40, 22, 12])
    safe_set_cell(ws, "B5", value="% Chi phí QL & BH (SG&A)")
    put(ws, 5, 4, P["sga"] / 100, F_PCT, kind="in")
    safe_set_cell(ws, "B6", value="% Thuế TNDN")
    put(ws, 6, 4, P["tax"] / 100, F_PCT, kind="in")
    lines = [
        ("Doanh thu thuần", "Tổng SL × FOB × Tỷ giá (B6)", f"={REVENUE}"),
        ("Chi phí NVL trực tiếp – TK 621", "Σ Q_MRP × Đơn giá", f"={q(21)}$G${tot21}"),
        ("Chi phí nhân công trực tiếp – TK 622", "Σ Tổng SL × Đơn giá NC", f"={q(22)}$E${tot22}"),
        ("Chi phí sản xuất chung – TK 627", "Σ Tổng SL × Đơn giá SXC", f"={q(23)}$E${tot22}"),
        ("Tổng giá vốn hàng bán (COGS)", "TK 621 + TK 622 + TK 627", "=D10+D11+D12"),
        ("Lợi nhuận gộp (Gross Profit)", "Doanh thu – COGS", "=D9-D13"),
        ("Chi phí QL & BH (OPEX)", "Doanh thu × % SG&A", "=D9*$D$5"),
        ("Lợi nhuận trước thuế & lãi vay (EBIT)", "Gross Profit – OPEX", "=D14-D15"),
        ("Thuế TNDN", "MAX(0, EBIT × % Thuế)", "=MAX(0,D16*$D$6)"),
        ("Lợi nhuận thuần (NPAT)", "MAX(0, EBIT × (1 – % Thuế))", "=MAX(0,D16*(1-$D$6))"),
    ]
    for s, (name, expl, f) in enumerate(lines):
        r = 9 + s
        bold = name.startswith(("Doanh thu", "Tổng giá vốn", "Lợi nhuận"))
        put(ws, r, 1, s + 1)
        put(ws, r, 2, name, bold=bold)
        put(ws, r, 3, expl)
        put(ws, r, 4, f, F_INT, kind="tot" if bold else None)
        put(ws, r, 5, f"=IF($D$9=0,0,D{r}/$D$9)", F_PCT)
    r = 9 + len(lines)
    for c in range(1, 6):
        X.put(ws, r, c, None, kind="tot").border = DBL_BORDER
    X.put(ws, r, 2, "TỔNG CHI PHÍ (COGS + OPEX + Thuế)", kind="tot").border = DBL_BORDER
    X.put(ws, r, 4, "=SUM(D13,D15,D17)", F_INT, kind="tot").border = DBL_BORDER
    X.put(ws, r, 5, f"=IF($D$9=0,0,D{r}/$D$9)", F_PCT, kind="tot").border = DBL_BORDER
    for ref in ("D14", "D18"):  # Lợi nhuận gộp, NPAT: xanh nếu dương, đỏ nếu âm
        ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=["0"],
                                      fill=PatternFill("solid", bgColor="C6EFCE"), font=Font(bold=True, color="006100")))
        ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0"],
                                      fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(bold=True, color="9C0006")))

    # ---------------- 25. BOM ERP đa cấp ----------------
    ws = X.sheet(25, ["STT", "Mã cha", "Loại cha", "Mã con", "Tên con", "Loại con", "ĐVT con", "Định mức",
                      "Hao hụt (%)", "Định mức gồm hao hụt", "Công đoạn"], [6, 30, 10, 34, 36, 10, 8, 14, 10, 16, 14],
                 subtitle="Rate(Parent, Child) × (1 + Scrap/100) – cấu trúc FG → SG → BTP → RM")
    for k, ((p, c), e) in enumerate(res["edges"].items()):
        r = 9 + k
        scrap = (e["eff"] / e["rate"] - 1) * 100 if e["rate"] else 0
        put(ws, r, 1, k + 1)
        put(ws, r, 2, p)
        put(ws, r, 3, nodes[p]["type"])
        put(ws, r, 4, c)
        put(ws, r, 5, nodes[c]["name"])
        put(ws, r, 6, nodes[c]["type"])
        put(ws, r, 7, nodes[c]["uom"])
        put(ws, r, 8, e["rate"], F_RATE, kind="norm")
        put(ws, r, 9, round(scrap, 6), F_NUM, kind="norm")
        put(ws, r, 10, f"=H{r}*(1+I{r}/100)", F_RATE)
        put(ws, r, 11, nodes[c]["stage"])

    return X.save()


def count_formulas(xlsx_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    stats = []
    for ws in wb.worksheets:
        n_f = sum(1 for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
        stats.append({"Sheet": ws.title, "Số ô công thức": n_f, "B6": ws["B6"].value if ws.title == SH[2] else ""})
    return wb.sheetnames, pd.DataFrame(stats)


# =============================================================================
# 9. GIAO DIỆN STREAMLIT
# =============================================================================

DEFAULT_STAGE_RATES = [  # (từ khóa, NC VND/SP, SXC VND/SP) – giá trị gợi ý, chỉnh tại Sidebar
    ("cat", 12000, 6000), ("chan", 8000, 4000), ("khoan", 6000, 3000), ("han", 25000, 12000),
    ("mai", 6000, 3000), ("son", 18000, 15000), ("ma", 15000, 12000), ("lap", 15000, 6000),
    ("dong goi", 8000, 3000),
]

MRP_FMT = {"Nhu cầu MRP": "{:,.2f}", "Đơn giá (VND)": "{:,.0f}", "Thành tiền (VND)": "{:,.0f}"}

ACCEPTANCE = {"R-O-S01-00-032-001": 862.50, "R-P-S01-052-00-002": 12540.00}

NUMCFG = st.column_config.NumberColumn


DEFAULT_STAGES = ["Cắt", "Chấn", "Hàn", "Mạ", "Đóng gói"]
TOT_LABEL = "TỔNG CỘNG"
TOT_STYLE = "font-weight:bold;background-color:#D9E1F2;border-top:2px double #1F3864"
GOOD_STYLE = "background-color:#C6EFCE;color:#006100;font-weight:bold"
BAD_STYLE = "background-color:#FFC7CE;color:#9C0006;font-weight:bold"


def parse_money(txt, default):
    """Đọc số có dấu phân cách hàng nghìn ('25,400' / '25 400'); dấu chấm là thập phân."""
    try:
        return float(re.sub(r"[,\s]", "", str(txt)))
    except ValueError:
        return float(default)


def money_input(label, key, default, help=None):
    """Ô nhập số hiển thị phân cách hàng nghìn (gõ tự do, tự định dạng lại khi rời ô)."""
    def _fmt():
        st.session_state[key] = f"{parse_money(st.session_state[key], default):,.0f}"
    if key not in st.session_state:
        st.session_state[key] = f"{default:,.0f}"
    return parse_money(st.text_input(label, key=key, on_change=_fmt, help=help), default)


def with_total(df, label_col, sum_cols, label=TOT_LABEL):
    """Thêm dòng TỔNG CỘNG ở cuối bảng."""
    t = {c: None for c in df.columns}
    t[label_col] = label
    for c in sum_cols:
        t[c] = float(pd.to_numeric(df[c], errors="coerce").sum())
    return pd.concat([df.astype(object), pd.DataFrame([t], dtype=object)], ignore_index=True)


def with_uom_totals(df):
    """Dòng tổng MRP tách theo từng ĐVT + dòng tổng thành tiền. Trả về (df, số dòng tổng)."""
    rows = []
    for u, g in df.groupby("ĐVT", sort=False):
        rows.append({"Tên NVL": uom_total_label(u), "ĐVT": u, "Nhu cầu MRP": float(g["Nhu cầu MRP"].sum()),
                     "Thành tiền (VND)": float(g["Thành tiền (VND)"].sum())})
    rows.append({"Tên NVL": "TỔNG THÀNH TIỀN (VND)", "Thành tiền (VND)": float(df["Thành tiền (VND)"].sum())})
    full = pd.concat([df.astype(object), pd.DataFrame(rows, dtype=object)], ignore_index=True)
    return full, len(rows)


def show_df(df, fmts, total=False, good_bad=(), good_bad_rows=None, **kw):
    """st.dataframe có phân cách hàng nghìn, dòng tổng in đậm, tô xanh/đỏ theo dấu (Conditional Formatting)."""
    sty = df.style.format(fmts, na_rep="")
    last = len(df) - 1

    def row_style(row):
        if total and row.name > last - int(total):
            return [TOT_STYLE] * len(row)
        return [""] * len(row)
    sty = sty.apply(row_style, axis=1)

    def sign_style(v):
        v = pd.to_numeric(v, errors="coerce")
        if pd.isna(v) or v == 0:
            return ""
        return GOOD_STYLE if v > 0 else BAD_STYLE
    if good_bad:
        sty = sty.map(sign_style, subset=list(good_bad))
    if good_bad_rows:  # (cột nhãn, từ khóa, cột giá trị) – tô theo dòng chỉ tiêu
        lab, keys, val = good_bad_rows
        mask = df[lab].astype(str).str.contains("|".join(keys), regex=True, na=False)
        sty = sty.apply(lambda c: [sign_style(v) if m else "" for v, m in zip(c, mask)], subset=[val])
    return st.dataframe(sty, hide_index=True, width="stretch", **kw)


def _default_rate(stage):
    n = norm(stage)
    for kw, lab, oh in DEFAULT_STAGE_RATES:
        if re.search(rf"\b{kw}", n):
            return lab, oh
    return 10000, 5000


def render_tree(res, fg, max_rows=3000):
    ch, nodes = res["children"], res["nodes"]
    out = []

    def go(n, d, rate, cumr, path):
        if len(out) >= max_rows:
            return
        node = nodes.get(n, {"name": "", "type": "", "uom": ""})
        out.append({"Cấp": d, "Cây cấu trúc": ("　" * d) + ("└─ " if d else "") + n, "Tên": node["name"],
                    "Loại": node["type"], "ĐVT": node["uom"], "Định mức cạnh (gồm HH)": rate,
                    "R_cum / 1 FG": cumr})
        for c, r in ch.get(n, []):
            if c not in path:
                go(c, d + 1, r, cumr * r, path | {c})

    go(fg, 0, 1.0, 1.0, frozenset([fg]))
    return pd.DataFrame(out)


def clean_slate():
    st.markdown(
        """
<div class="hero">
  <h2>🏭 Hệ thống đang ở trạng thái chờ (Clean Slate)</h2>
  <p>Vui lòng nạp <b>tối thiểu 2 tệp bắt buộc</b> ở phía trên để kích hoạt bộ máy rã BOM & MRP:</p>
</div>""", unsafe_allow_html=True)
    c1, c2, c3 = st.columns(3)
    c1.markdown("##### 1️⃣ Sales Order (SO)\n**BẮT BUỘC** – Sản lượng & Giá bán FOB ($USD) theo mã SKU thành phẩm.")
    c2.markdown("##### 2️⃣ Technical BOM\n**BẮT BUỘC** – Cấu trúc kỹ thuật đa cấp, tiêu hao NVL & công đoạn gia công.")
    c3.markdown("##### 3️⃣ Purchase Order (PO)\n*TÙY CHỌN* – Đơn giá mua thực tế, số PO & nhà cung cấp.")
    st.info("📌 Kịch bản 2 tệp (SO + BOM): mở khóa 100% tính toán, đơn giá lấy từ BOM / Live Editor / Sidebar.  \n"
            "📌 Kịch bản 3 tệp (SO + BOM + PO): đối chiếu danh mục NVL rã từ BOM sang PO để lấy đơn giá & số PO thực tế.")


def main():
    st.set_page_config(page_title="METALIC ERP Costing & MRP", page_icon="🏭", layout="wide")
    st.markdown("""
<style>
.hero{padding:1.2rem 1.5rem;border-radius:12px;background:linear-gradient(90deg,#1F3864,#2E75B6);color:#fff;margin-bottom:1rem}
.hero h2{color:#fff;margin:0 0 .4rem 0}
div[data-testid="stMetricValue"]{font-size:1.35rem}
header[data-testid="stHeader"]{display:none !important}
.block-container{padding-top:.5rem !important;padding-bottom:.5rem !important}
[data-testid="stSidebar"]{padding-top:.5rem !important}
.app-title{font-size:21px;font-weight:700;margin:0;padding:0;line-height:1.6}
.app-badge{font-size:.8rem;text-align:right;padding-top:.35rem;color:#2E7D32;font-weight:600}
.app-badge.wait{color:#B26A00}
[data-testid="stHorizontalBlock"]{margin-bottom:0 !important}
[data-testid="stVerticalBlock"]{gap:.5rem}
[data-testid="stSidebarHeader"]{padding-top:0 !important;height:1.5rem !important}
[data-testid="stSidebarUserContent"]{padding-top:.5rem !important}
[data-testid="stFileUploaderDropzone"]{padding:.35rem .6rem !important;min-height:0 !important}
[data-testid="stFileUploaderDropzoneInstructions"]{display:none !important}
[data-testid="stFileUploader"] label p{font-size:.85rem;margin-bottom:0}
[data-testid="stFileUploader"] section+div{padding-top:.2rem}
.btn-spacer{height:1.65rem}
.st-key-act_btns button{font-size:13px !important;font-weight:600 !important;padding:.35rem .5rem !important;
white-space:nowrap !important;line-height:1.2 !important;margin-bottom:.25rem !important}
</style>""", unsafe_allow_html=True)
    h_title, h_status = st.columns([6, 4], vertical_alignment="center")
    h_title.markdown(f'<p class="app-title">🏭 {APP_TITLE}</p>', unsafe_allow_html=True)
    status_ph = h_status.empty()

    # ---------------- Nạp dữ liệu đầu vào (màn hình chính) ----------------
    gen = st.session_state.get("_gen", 0)  # đổi số này để reset các ô file uploader
    up1, up2, up3, b1 = st.columns([2.5, 2.5, 2.5, 2.2], vertical_alignment="top")
    with up1:
        so_file = st.file_uploader("1. SO (BẮT BUỘC)", type=["xlsx", "xlsm", "csv"], key=f"up_so_{gen}")
    with up2:
        bom_file = st.file_uploader("2. BOM (BẮT BUỘC)", type=["xlsx", "xlsm", "csv"], key=f"up_bom_{gen}")
    with up3:
        po_file = st.file_uploader("3. PO (TÙY CHỌN)", type=["xlsx", "xlsm", "csv"], key=f"up_po_{gen}")
    with b1:
        st.markdown('<div class="btn-spacer"></div>', unsafe_allow_html=True)
        with st.container(key="act_btns"):
            btn_calc = st.button("🚀 TÍNH TOÁN MRP", type="primary", width="stretch",
                                 disabled=not (so_file and bom_file),
                                 help="Chỉ khi nhấn nút này hệ thống mới đọc file, rã BOM, tính MRP và giá thành.")
            btn_reset = st.button("🗑️ XÓA DỮ LIỆU", width="stretch",
                                  help="Xóa toàn bộ kết quả đã tính và các tệp đã nạp, quay về trạng thái chờ.")

    with st.sidebar:
        st.header("⚙️ Thông số hệ thống")
        P = {
            "project": st.text_input("Tên dự án / Khách hàng", "METALIC EXPORT PROJECT"),
            "er": max(1.0, money_input("Tỷ giá USD/VND (ô B6)", "in_er", 25400)),
            "start_date": st.date_input("Ngày bắt đầu sản xuất", dt.date.today()),
            "lead_days": st.number_input("Lead time sản xuất (ngày)", min_value=0, value=21, step=1),
            "sga": st.number_input("% Chi phí QL & BH (SG&A)", min_value=0.0, max_value=100.0, value=5.0, step=0.5),
            "tax": st.number_input("% Thuế TNDN", min_value=0.0, max_value=100.0, value=20.0, step=1.0),
            "default_scrap": st.number_input("% Hao hụt mặc định (khi BOM trống)", min_value=0.0, value=0.0, step=0.5),
            "kg_price": max(0.0, money_input("Giá tham khảo NVL (VND/Kg)", "in_kg", 22000)),
            "unit_price": max(0.0, money_input("Giá tham khảo vật tư (ĐVT Cái/Khác, VND)", "in_unit", 5000)),
            "apply_sidebar": st.checkbox("Áp giá tạm từ Sidebar cho NVL chưa có giá", value=False,
                                         help="Mặc định tắt: NVL không có giá PO thì đơn giá = 0đ. Bật để dùng 2 ô giá tham khảo ở trên."),
            "rm_prefix": st.text_input("Tiền tố mã NVL thô trong BOM", "R-"),
            "pallet_qty": st.number_input("SL thành phẩm / Pallet", min_value=1, value=50, step=1),
        }
        rate_box = st.container()  # bảng đơn giá công đoạn – điền sau (mặc định khi chưa có BOM)

    def render_rates(stage_names, key):
        with rate_box:
            st.header("🧮 Đơn giá công đoạn (TK 622 / 627)")
            rate_df = pd.DataFrame([{"Công đoạn": s, "Nhân công VND/SP": float(_default_rate(s)[0]),
                                     "SXC VND/SP": float(_default_rate(s)[1])} for s in stage_names],
                                   columns=["Công đoạn", "Nhân công VND/SP", "SXC VND/SP"])
            tip = "Gõ trực tiếp hoặc dùng mũi tên ↑/↓ (bước 1,000)"
            ed = st.data_editor(rate_df, key=key, hide_index=True, width="stretch", disabled=["Công đoạn"],
                                column_config={
                                    "Nhân công VND/SP": NUMCFG(format="localized", min_value=0, step=1000, help=tip),
                                    "SXC VND/SP": NUMCFG(format="localized", min_value=0, step=1000, help=tip)})
            st.caption("Danh mục công đoạn lấy động từ Technical BOM đã nạp; chưa có BOM thì dùng danh mục mặc định. "
                       "Giá trị gợi ý – hãy điều chỉnh theo định mức nhà máy.")
        return ed

    # ---------------- Nút XÓA DỮ LIỆU ----------------
    if btn_reset:
        gen_next = st.session_state.get("_gen", 0) + 1
        st.session_state.clear()
        st.session_state["_gen"] = gen_next  # key uploader mới → các ô tải file trống
        st.rerun()

    # ---------------- Nút TÍNH TOÁN MRP: đọc file 1 lần, lưu vào session_state ----------------
    cur_sig = None
    if so_file and bom_file:
        cur_sig = (file_hash(so_file.getvalue()), file_hash(bom_file.getvalue()),
                   file_hash(po_file.getvalue()) if po_file else None)
    if btn_calc and cur_sig:
        st.session_state.pop("calc", None)
        try:
            so_df, so_warn, so_info = parse_so(so_file.getvalue(), so_file.name)
            bom_df, stage_cols, sheet_ctx, bom_infos, bom_warn = parse_bom(bom_file.getvalue(), bom_file.name)
        except Exception as e:
            st.error(f"❌ Lỗi đọc file: {e}")
            st.stop()
        po_df, po_info, po_msg = None, None, None
        if po_file:
            try:
                po_df, _, po_info = parse_po(po_file.getvalue(), po_file.name)
            except Exception as e:
                po_msg = f"⚠️ Không đọc được file PO ({e}) – chuyển sang kịch bản 2 tệp."
                po_df = None
        st.session_state["calc"] = {
            "sig": cur_sig, "so_df": so_df, "so_warn": so_warn, "so_info": so_info,
            "bom_df": bom_df, "stage_cols": stage_cols, "sheet_ctx": sheet_ctx,
            "bom_infos": bom_infos, "bom_warn": bom_warn, "po_df": po_df, "po_info": po_info,
            "po_msg": po_msg, "bom_key": cur_sig[1],
        }

    calc = st.session_state.get("calc")
    if not calc:
        render_rates(DEFAULT_STAGES, "rates_default")
        clean_slate()
        if cur_sig:
            st.warning("📂 Đã nạp đủ tệp. Nhấn **🚀 TÍNH TOÁN MRP** ở phía trên để bắt đầu.")
        st.stop()
    if calc["sig"] != cur_sig:
        render_rates([c[len(STAGE_PREFIX):] for c in calc["stage_cols"]], f"rates_{calc['bom_key']}")
        st.warning("⚠️ Tệp đầu vào đã thay đổi so với lần tính trước. Nhấn **🚀 TÍNH TOÁN MRP** để tính lại "
                   "hoặc **🗑️ XÓA DỮ LIỆU** để làm lại từ đầu.")
        st.stop()

    so_df, so_warn, so_info = calc["so_df"], calc["so_warn"], calc["so_info"]
    bom_df, stage_cols, sheet_ctx = calc["bom_df"], calc["stage_cols"], calc["sheet_ctx"]
    bom_infos, bom_warn = calc["bom_infos"], calc["bom_warn"]
    po_df, po_info = calc["po_df"], calc["po_info"]
    if calc["po_msg"]:
        st.warning(calc["po_msg"])
    scenario = "3 tệp (SO + BOM + PO)" if po_df is not None else "2 tệp (SO + BOM)"
    bom_key = calc["bom_key"]

    tabs = st.tabs(["📊 Tổng quan", "🧾 Đơn hàng SO & PO", "🛠️ Live BOM Editor", "🌳 BOM ERP đa cấp",
                    "📦 Hoạch định MRP", "💰 Giá thành & P&L", "📥 Xuất Excel Master"])

    # ---------------- Live BOM Editor ----------------
    with tabs[2]:
        st.subheader("🛠️ Live BOM Editor – chỉnh sửa trực tiếp Technical BOM")
        st.caption("Mọi thay đổi (SL, KL NVL, hao hụt, đơn giá, thứ tự công đoạn, Cấp) được rã lại ngay lập tức. "
                   "Cây cha–con dựng theo thứ tự dòng + cột **Cấp** (suy ra từ MỤC 1 / 1.1 / 1.1.1 hoặc khoảng lùi đầu dòng).")
        cfg = {c: NUMCFG(c, format="%.4g") for c in ["Dày", "Rộng", "Dài", "SL/Cha", "KL NVL (kg)", "Hao hụt (%)"]}
        cfg["Cấp"] = NUMCFG("Cấp", min_value=0, step=1, format="%d")
        cfg["Đơn giá"] = NUMCFG("Đơn giá", format="%.0f")
        cfg["Dòng"] = NUMCFG("Dòng", format="%d", disabled=True, help="Dòng vật lý trong file Excel gốc")
        cfg["Sheet"] = st.column_config.TextColumn("Sheet", disabled=True, help="Tên Sheet nguồn")
        for c in stage_cols:
            cfg[c] = NUMCFG(c[len(STAGE_PREFIX):], min_value=0, step=1, format="%g",
                            help="Thứ tự công đoạn tại dòng (1, 2, 3...). Để 0 hoặc trống = không thực hiện.")
        edited = st.data_editor(bom_df, key=f"bom_editor_{bom_key}", num_rows="dynamic", width="stretch",
                                height=420, column_config=cfg)
        with st.expander("🔎 Kết quả Smart Header Parser cho Technical BOM"):
            for inf in bom_infos:
                st.markdown(f"**Sheet `{inf['sheet']}`** – {inf['status']}")
                if inf.get("mapping"):
                    st.json({"cột": inf["mapping"], "công đoạn": inf.get("stages", [])}, expanded=False)

    # ---------------- Bảng đơn giá công đoạn (Sidebar) ----------------
    stage_names = [c[len(STAGE_PREFIX):] for c in stage_cols]
    rate_ed = render_rates(stage_names, f"rates_{bom_key}")

    # ---------------- Bộ máy tính toán ----------------
    res = run_engine(edited, stage_cols, sheet_ctx, so_df, po_df, P)
    used = res["used_stages"]
    rate_map = {r["Công đoạn"]: r for r in rate_ed.to_dict("records")}
    stage_rates = [{"stage": s, "labor": to_float(rate_map.get(s, {}).get("Nhân công VND/SP"), 0),
                    "oh": to_float(rate_map.get(s, {}).get("SXC VND/SP"), 0)} for s in used]
    fin = compute_financials(res, so_df, stage_rates, P)
    mrp = res["mrp_df"]
    nodes = res["nodes"]
    all_warn = so_warn + bom_warn + res["warnings"]

    # ---------------- Tổng quan ----------------
    with tabs[0]:
        status_ph.markdown(f'<div class="app-badge">✅ Kịch bản {scenario} – bộ máy MRP đã kích hoạt</div>', unsafe_allow_html=True)
        c = st.columns(4)
        c[0].metric("Dòng SO / Tổng SL", f"{len(so_df)} / {fin['total_qty']:,.0f}")
        c[1].metric("Mã NVL thô (RM)", f"{len(mrp)}")
        c[2].metric("Mã BTP sinh tự động", f"{sum(1 for n in nodes.values() if n['type'] == 'BTP')}")
        c[3].metric("Cạnh BOM ERP", f"{len(res['edges'])}")
        c = st.columns(4)
        c[0].metric("Doanh thu (VND)", f"{fin['revenue']:,.0f}")
        c[1].metric("COGS (VND)", f"{fin['cogs']:,.0f}")
        c[2].metric("Lợi nhuận gộp", f"{fin['gp']:,.0f}",
                    f"{fin['gp'] / fin['revenue'] * 100:.1f}% DT" if fin["revenue"] else None)
        c[3].metric("NPAT (VND)", f"{fin['npat']:,.0f}")
        found = [k for k in ACCEPTANCE if k in set(mrp["Mã NVL"])]
        if found:
            st.markdown("##### ✅ Kiểm tra nghiệm thu MRP")
            chk = []
            for k in found:
                v = float(mrp.loc[mrp["Mã NVL"] == k, "Nhu cầu MRP"].iloc[0])
                chk.append({"Mã NVL": k, "MRP tính được": v, "Số liệu phân xưởng": ACCEPTANCE[k],
                            "Kết quả": "✅ Khớp" if abs(v - ACCEPTANCE[k]) < 0.005 else "❌ Lệch"})
            st.dataframe(pd.DataFrame(chk), hide_index=True, width="stretch",
                         column_config={"MRP tính được": NUMCFG(format="%.2f"), "Số liệu phân xưởng": NUMCFG(format="%.2f")})
        h1, h2 = st.columns([3, 1], vertical_alignment="center")
        h1.markdown("##### 🏷️ NVL theo giá trị")
        top_sel = h2.selectbox("Hiển thị", ["Top 5", "Top 10", "Top 25", "Tất cả NVL"], index=1, key="top_n",
                               label_visibility="collapsed")
        if len(mrp):
            n_top = {"Top 5": 5, "Top 10": 10, "Top 25": 25}.get(top_sel, len(mrp))
            show_df(mrp.sort_values("Thành tiền (VND)", ascending=False).head(n_top), MRP_FMT, height=300)
        # Ẩn danh sách cảnh báo dữ liệu để giữ Dashboard gọn (mảng all_warn vẫn được tính ở backend)
        # if all_warn:
        #     with st.expander(f"⚠️ Cảnh báo dữ liệu ({len(all_warn)})"):
        #         for w in all_warn:
        #             st.write("• " + w)

    # ---------------- SO & PO ----------------
    with tabs[1]:
        t1, t2 = st.tabs(["Sales Order (chuẩn hóa SO_SCHEMA)", "Purchase Order (chuẩn hóa PO_SCHEMA)"])
        with t1:
            st.caption(f"Sheet `{so_info['sheet']}` • tiêu đề tại dòng {so_info['header_row']}")
            so_v = so_df.copy()
            so_v["Thành tiền ($USD)"] = so_v["Số Lượng"] * so_v["Giá Bán FOB ($USD)"]
            so_v["Thành tiền (VND)"] = so_v["Thành tiền ($USD)"] * P["er"]
            so_v = with_total(so_v, "Tên Thành Phẩm", ["Số Lượng", "Thành tiền ($USD)", "Thành tiền (VND)"])
            show_df(so_v, {"STT": "{:.0f}", "Số Lượng": "{:,.0f}", "Giá Bán FOB ($USD)": "{:,.2f}",
                           "Thành tiền ($USD)": "{:,.2f}", "Thành tiền (VND)": "{:,.0f}"}, total=True, height=340)
            with st.expander("Ánh xạ cột (Smart Header Parser)"):
                st.json(so_info["mapping"])
        with t2:
            if po_df is None:
                st.info("Chưa nạp file PO – kịch bản 2 tệp. Đơn giá NVL = 0đ, trừ khi nhập giá tại Live BOM Editor.")
            else:
                st.caption(f"Sheet `{po_info['sheet']}` • tiêu đề tại dòng {po_info['header_row']}")
                po_v = po_df.rename(columns={"Số PO Mua": "Số PO"})
                po_v["Thành tiền dự kiến (VND)"] = po_v["Số Lượng Mua Thực Tế"] * po_v["Đơn Giá Mua (VND)"]
                po_v = with_total(po_v, "Tên Hàng Hóa Phụ Kiện Mua", ["Số Lượng Mua Thực Tế", "Thành tiền dự kiến (VND)"])
                show_df(po_v, {"Số Lượng Mua Thực Tế": "{:,.2f}", "Đơn Giá Mua (VND)": "{:,.0f}",
                               "Thành tiền dự kiến (VND)": "{:,.0f}"}, total=True, height=340)
                with st.expander("Ánh xạ cột (Smart Header Parser)"):
                    st.json(po_info["mapping"])
                matched = int(mrp["Nguồn giá"].str.startswith("PO:").sum()) if len(mrp) else 0
                st.metric("NVL khớp đơn giá PO", f"{matched} / {len(mrp)}")

    # ---------------- BOM ERP ----------------
    with tabs[3]:
        t1, t2, t3, t4 = st.tabs(["🌳 Cây cấu trúc theo FG", "🔗 Danh sách cạnh BOM ERP", "📇 Danh mục mã (Item Master)",
                                  "⚙️ Mã BTP công đoạn"])
        with t1:
            fgs = list(dict.fromkeys(l["Mã Sản Phẩm ERP"] for l in res["so_lines"]))
            fg = st.selectbox("Chọn thành phẩm", fgs)
            tree = render_tree(res, fg)
            st.dataframe(tree, hide_index=True, width="stretch", height=400,
                         column_config={"Định mức cạnh (gồm HH)": NUMCFG(format="%.6g"), "R_cum / 1 FG": NUMCFG(format="%.6g")})
        with t2:
            e_df = pd.DataFrame([{"Mã cha": p, "Loại cha": nodes[p]["type"], "Mã con": c, "Tên con": nodes[c]["name"],
                                  "Mã NVL / Quy cách": (f"{c} | {nodes[c]['std'] or nodes[c]['name']}"
                                                        if nodes[c]["type"] == "RM" else ""),
                                  "Tên NVL (cấp lá)": ((nodes[c]["std"] or nodes[c]["name"])
                                                       if nodes[c]["type"] == "RM" else ""),
                                  "Loại con": nodes[c]["type"], "ĐVT": nodes[c]["uom"], "Định mức": e["rate"],
                                  "Hao hụt (%)": (e["eff"] / e["rate"] - 1) * 100 if e["rate"] else 0,
                                  "Định mức gồm HH": e["eff"]} for (p, c), e in res["edges"].items()])
            st.dataframe(e_df, hide_index=True, width="stretch", height=400,
                         column_config={"Định mức": NUMCFG(format="%.6g"), "Hao hụt (%)": NUMCFG(format="%.2f"),
                                        "Định mức gồm HH": NUMCFG(format="%.6g")})
        with t3:
            n_df = pd.DataFrame([{"Mã": n["code"], "Tên kỹ thuật": n["name"], "Loại": TYPE_LABEL.get(n["type"], n["type"]),
                                  "ĐVT": n["uom"], "Tên NVL chuẩn hóa": n["std"], "Công đoạn": n["stage"]} for n in nodes.values()])
            st.dataframe(n_df, hide_index=True, width="stretch", height=400)
        with t4:
            b_df = pd.DataFrame([{"Mã BTP": n["code"], "Công đoạn": n["stage"], "Tên": n["name"]}
                                 for n in nodes.values() if n["type"] == "BTP"], columns=["Mã BTP", "Công đoạn", "Tên"])
            st.caption("Cú pháp: {Mã_Cha}.BTP_{Tên_Công_Đoạn} – mỗi công đoạn được đánh dấu tại dòng Mã Cha sinh 1 mã WIP.")
            st.dataframe(b_df, hide_index=True, width="stretch", height=400)

    # ---------------- MRP ----------------
    with tabs[4]:
        t1, t2, t3 = st.tabs(["📦 Tổng hợp nhu cầu MRP", "🧮 Ma trận R_cum (RM × FG)", "🗓️ Lịch rã WIP (DPP)"])
        with t1:
            mrp_t, n_tot = with_uom_totals(mrp)
            show_df(mrp_t, MRP_FMT, total=n_tot, height=400)
            mrp_buf = io.BytesIO()
            with pd.ExcelWriter(mrp_buf, engine="openpyxl") as writer:
                mrp.to_excel(writer, index=False, sheet_name="Hoach_Dinh_MRP")
            st.download_button("⬇️ Tải bảng MRP (Excel)", mrp_buf.getvalue(), "Metalic_MRP_Export.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        with t2:
            st.latex(r"R_{cum}(RM_k,FG_i)=\sum_{P}\prod_{(Pa,Ch)\in P} Rate(Pa,Ch)\times\left(1+\frac{Scrap(Pa,Ch)}{100}\right)")
            cols = [f"{l['Mã Sản Phẩm ERP']} (#{j + 1})" for j, l in enumerate(res["so_lines"])]
            mat = pd.DataFrame(res["R"], columns=cols)
            mat.insert(0, "Mã NVL", res["rm_used"])
            mat.insert(1, "ĐVT", [nodes[c]["uom"] for c in res["rm_used"]])
            st.dataframe(mat, hide_index=True, width="stretch", height=400,
                         column_config={c: NUMCFG(format="%.6g") for c in cols})
        with t3:
            wdf = pd.DataFrame([{"Mã TP": res["so_lines"][w["j"]]["Mã Sản Phẩm ERP"], "Mã BTP / Chi tiết": w["code"],
                                 "Tên": w["name"], "Công đoạn": w["stage"], "Định mức tích lũy": w["rate"],
                                 "SL kế hoạch": res["so_lines"][w["j"]]["Số Lượng"] * w["rate"],
                                 "Ngày bắt đầu": P["start_date"] + dt.timedelta(days=w["offset"])} for w in res["wip"]])
            st.dataframe(wdf, hide_index=True, width="stretch", height=400)

    # ---------------- Giá thành & P&L ----------------
    with tabs[5]:
        t1, t2, t3, t4, t5 = st.tabs(["🏷️ Giá thành SKU", "TK 621 – NVL", "TK 622 – Nhân công", "TK 627 – SXC",
                                      "📈 Báo cáo P&L"])
        money = NUMCFG(format="localized")
        if len(mrp) and fin["tk621"] <= 0:
            st.warning("Dự án đang tính toán với Đơn giá NVL = 0đ do chưa nạp file PO hoặc chưa nhập giá trong Live BOM Editor. "
                       "Vui lòng nạp PO để hạch toán đầy đủ TK 621 và Giá vốn COGS.")
        with t1:
            sku_t = with_total(fin["sku"], "Tên TP", ["Số lượng", "Tổng giá thành"])
            show_df(sku_t, {k: "{:,.0f}" for k in ["Số lượng", "CP NVL/SP", "CP NC/SP", "CP SXC/SP", "Giá thành/SP",
                                                  "Tổng giá thành", "Giá bán VND/SP", "Lãi gộp/SP"]}
                    | {"Biên LN gộp": "{:.1%}"}, total=True, good_bad=["Lãi gộp/SP", "Biên LN gộp"])
        with t2:
            st.metric("Tổng TK 621", f"{fin['tk621']:,.0f} VND")
            t621, n_tot = with_uom_totals(mrp[["Mã NVL", "Tên NVL", "ĐVT", "Nhu cầu MRP", "Đơn giá (VND)",
                                               "Thành tiền (VND)", "Nguồn giá"]])
            show_df(t621, MRP_FMT, total=n_tot)
        for tab, field, total, label in ((t3, "labor", fin["tk622"], "TK 622"), (t4, "oh", fin["tk627"], "TK 627")):
            with tab:
                st.metric(f"Tổng {label}", f"{total:,.0f} VND")
                d = pd.DataFrame([{"Công đoạn": r["stage"], "Đơn giá / SP": r[field], "Tổng SL SO": fin["total_qty"],
                                   "Thành tiền (VND)": r[field] * fin["total_qty"]} for r in stage_rates],
                                 columns=["Công đoạn", "Đơn giá / SP", "Tổng SL SO", "Thành tiền (VND)"])
                d = with_total(d, "Công đoạn", ["Đơn giá / SP", "Thành tiền (VND)"])
                show_df(d, {"Đơn giá / SP": "{:,.0f}", "Tổng SL SO": "{:,.0f}", "Thành tiền (VND)": "{:,.0f}"}, total=True)
                if not stage_rates:
                    st.info("BOM chưa đánh dấu công đoạn nào.")
        with t5:
            pl = fin["pl"].copy()
            pl.loc[len(pl)] = ["TỔNG CHI PHÍ (COGS + OPEX + Thuế)", fin["cogs"] + fin["opex"] + fin["tax"]]
            show_df(pl, {"Giá trị (VND)": "{:,.0f}"}, total=True,
                    good_bad_rows=("Chỉ tiêu", ["Lợi nhuận gộp", "NPAT"], "Giá trị (VND)"))
            if fin["ebit"] < 0:
                st.warning("EBIT âm – theo đặc tả NPAT = MAX(0, EBIT × (1 − %Thuế)) nên NPAT hiển thị 0.")

    # ---------------- Xuất Excel ----------------
    with tabs[6]:
        st.subheader("📥 Excel Master 25 Sheet – 100% công thức liên kết sống")
        try:
            xbytes = build_master_excel(res, fin, so_df, stage_rates, P)
            names, stats = count_formulas(xbytes)
            c = st.columns(3)
            c[0].metric("Số sheet", len(names))
            c[1].metric("Tổng ô công thức", f"{int(stats['Số ô công thức'].sum()):,}")
            c[2].metric("Ô B6 (Tỷ giá)", f"{P['er']:,.0f}")
            st.download_button("⬇️ Tải METALIC_ERP_MASTER.xlsx", xbytes,
                               f"METALIC_ERP_MASTER_{dt.date.today():%Y%m%d}.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
            st.dataframe(stats[["Sheet", "Số ô công thức"]], hide_index=True, width="stretch", height=400)
        except Exception as e:
            st.error(f"❌ Lỗi kết xuất Excel: {e}")


if __name__ == "__main__":
    main()
