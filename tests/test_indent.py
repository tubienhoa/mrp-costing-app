import os
import sys, io, datetime as dt, warnings, logging
warnings.filterwarnings("ignore"); logging.disable(logging.CRITICAL)
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import app, openpyxl
from openpyxl.styles import Alignment

so_csv = "STT;Mã SP KH;Mã SP ERP;Tên TP;ĐVT;Số lượng;Giá FOB (USD)\n1;K1;FG-RACK;Kệ thép;Cái;10;120,5\n".encode("utf-8-sig")
so_df, w, _ = app.parse_so(so_csv, "so.csv")
print(so_df.to_string())

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "RACK"
for i, h in enumerate(["STT", "Mã", "Mô tả", "Dày", "Rộng", "Dài", "SL", "KL nguyên liệu (kg)", "Hàn", "Sơn"], 1):
    ws.cell(1, i, h)
rows = [(1, "FG-RACK", "Kệ thép", 0, None, None, None, 1, None, None, "x"),
        (2, "KHUNG", "Khung kệ", 1, None, None, None, 2, None, "x", None),
        (3, "", "Thép tròn đặc", 2, None, 16, 500, 4, 0.79, None, None),
        (4, "", "Tấm đáy thép tấm", 2, 3, 400, 400, 1, 3.77, None, None)]
for r, (stt, code, desc, ind, t, w_, l, qty, kg, han, son) in enumerate(rows, 2):
    for c, v in enumerate([stt, code, desc, t, w_, l, qty, kg, han, son], 1):
        if v is not None: ws.cell(r, c, v)
    ws.cell(r, 3).alignment = Alignment(indent=ind)
b = io.BytesIO(); wb.save(b)
bom_df, sc, ctx, infos, _ = app.parse_bom(b.getvalue(), "bom.xlsx")
print(bom_df[["MỤC", "Cấp", "Mã chi tiết", "Mô tả"] + sc].to_string())
P = dict(project="T", er=25400.0, start_date=dt.date(2026, 10, 2), lead_days=21, sga=5.0, tax=20.0, default_scrap=0.0,
         kg_price=22000.0, unit_price=0.0, rm_prefix="R-", pallet_qty=50)
res = app.run_engine(bom_df, sc, ctx, so_df, None, P)
print(res["mrp_df"][["Mã NVL", "Tên NVL", "ĐVT", "Nhu cầu MRP"]].to_string())
print([c for c, n in res["nodes"].items() if n["type"] == "BTP"], res["warnings"])
